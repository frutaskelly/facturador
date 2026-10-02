"""SKU del cliente capturado desde la lista de precios (dueño, 2-oct-2026).

Balles y Jubran comparten una lista y necesitan SU SKU en el XML. La lista lo
captura —uno por uno o en su Excel— y lo guarda en el catálogo de CADA cliente
asignado, que es lo que leen el XML del Facturador y el masivo de SAE.
"""
import io
import uuid

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from app.core.auth import Principal, get_principal
from app.core.db import SessionLocal
from app.main import app
from app.models import (ClaveSae, Cliente, ClienteExterno, ListaAsignacion, ListaPrecios, Membership, Precio,
                        Producto, ProductoCliente, Role, Tenant, User)


@pytest.fixture
def env(db_engine):
    suffix = uuid.uuid4().hex[:8]
    db = SessionLocal()
    created = {"memberships": [], "users": [], "tenants": []}
    try:
        t = Tenant(slug=f"sku-{suffix}", legal_name="Sku SA", rfc=f"SK{suffix.upper()}"[:13],
                   regimen_fiscal_sat="601", domicilio_fiscal_cp="44100", tier="PRINCIPAL", status="ACTIVE")
        db.add(t); db.flush(); created["tenants"].append(t.id)
        role = db.query(Role).filter(Role.nombre == "ADMIN", Role.es_preset.is_(True)).one()
        sub = f"sub-sku-{suffix}"
        u = User(email=f"sku-{suffix}@t.test", auth_user_id=sub, full_name="admin")
        db.add(u); db.flush(); created["users"].append(u.id)
        m = Membership(tenant_id=t.id, user_id=u.id, role_id=role.id)
        db.add(m); db.flush(); created["memberships"].append(m.id)

        balles = Cliente(tenant_id=t.id, legal_name="OPERADORA BALLES VEGA", rfc="OBV010101AAA")
        jubran = Cliente(tenant_id=t.id, legal_name="DISTRIBUIDORA JUBRAN", rfc="DJU010101AAA")
        prod = Producto(tenant_id=t.id, sku="00000403", nombre="ZANAHORIA", unidad_base="KILO",
                        presentaciones={"KILO": 1, "PIEZA": {"factor": 1, "clave_sae": "ZANAHORIAPZ"}},
                        clave_sae="ZANAHORIAKG", clave_sat="50404100", unidad_sat="KGM")
        lista = ListaPrecios(tenant_id=t.id, codigo="BJ", nombre="Balles y Jubran")
        db.add_all([balles, jubran, prod, lista]); db.flush()
        for c in (balles, jubran):
            db.add(ListaAsignacion(tenant_id=t.id, lista_id=lista.id, cliente_id=c.id))
        db.add(ClienteExterno(tenant_id=t.id, sistema="SAE", clave="02:7", clave_normalizada="02:7",
                              cliente_id=balles.id))
        db.add(ClaveSae(tenant_id=t.id, empresa="02", clave="ZANA-FRUT-508", descripcion="ZANAHORIA", activa=True))
        for pres, precio in (("KILO", 16), ("PIEZA", 5)):
            db.add(Precio(tenant_id=t.id, lista_id=lista.id, producto_id=prod.id, presentacion=pres,
                          precio_unitario=precio))
        db.commit()
        yield {"admin": {"sub": sub, "email": u.email, "tenant_id": t.id}, "tenant_id": t.id,
               "lista": str(lista.id), "prod": str(prod.id), "balles": balles.id, "jubran": jubran.id}
    finally:
        for table in ("producto_clientes", "precios", "lista_asignaciones", "listas_precios", "claves_sae",
                      "cliente_externos", "productos", "clientes"):
            for tid in created["tenants"]:
                db.execute(text(f"DELETE FROM {table} WHERE tenant_id = :tid"), {"tid": tid})
        for mid in created["memberships"]:
            db.query(Membership).filter(Membership.id == mid).delete()
        for uid in created["users"]:
            db.query(User).filter(User.id == uid).delete()
        for tid in created["tenants"]:
            db.query(Tenant).filter(Tenant.id == tid).delete()
        db.commit(); db.close()


@pytest.fixture
def auth_as():
    def _set(user):
        app.dependency_overrides[get_principal] = lambda: Principal(
            auth_user_id=user["sub"], email=user["email"], role="authenticated", claims={"sub": user["sub"]})
    yield _set
    app.dependency_overrides.pop(get_principal, None)


def _h(env):
    return {"X-Tenant-Id": str(env["tenant_id"])}


def _filas_cliente(env):
    with SessionLocal() as s:
        return {(r.cliente_id, r.codigo_cliente, r.presentacion)
                for r in s.query(ProductoCliente).filter(ProductoCliente.tenant_id == env["tenant_id"])}


def _activar(client, env):
    r = client.patch(f"/api/v1/listas-precios/{env['lista']}", headers=_h(env), json={"lleva_sku_cliente": True})
    assert r.status_code == 200 and r.json()["lleva_sku_cliente"] is True


def test_el_sku_se_guarda_en_el_catalogo_de_cada_cliente_de_la_lista(client, env, auth_as):
    auth_as(env["admin"]); h = _h(env)
    url = f"/api/v1/listas-precios/{env['lista']}/sku-cliente"
    body = {"producto_id": env["prod"], "presentacion": "KILO", "sku": " zana-frut-508 "}
    assert client.put(url, headers=h, json=body).status_code == 409      # la lista no lo lleva aún
    _activar(client, env)
    r = client.put(url, headers=h, json=body)
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["sku"] == "ZANA-FRUT-508" and not out["distintos"] and out["avisos"] == []
    assert _filas_cliente(env) == {(env["balles"], "ZANA-FRUT-508", None), (env["jubran"], "ZANA-FRUT-508", None)}

    lista = client.get(url, headers=h).json()
    assert {c["nombre"] for c in lista["clientes"]} == {"OPERADORA BALLES VEGA", "DISTRIBUIDORA JUBRAN"}
    assert {(x["presentacion"], x["sku"]) for x in lista["renglones"]} == {("KILO", "ZANA-FRUT-508"), ("PIEZA", None)}

    # Un SKU por producto y cliente: el de PIEZA no pisa en silencio al de KILO.
    r = client.put(url, headers=h, json={"producto_id": env["prod"], "presentacion": "PIEZA", "sku": "ZANAPZ"})
    assert r.status_code == 409 and "KILO" in r.text

    # Vacío lo quita: el XML vuelve a llevar el SKU interno.
    r = client.put(url, headers=h, json={"producto_id": env["prod"], "presentacion": "KILO", "sku": ""})
    assert r.status_code == 200 and r.json()["sku"] is None
    assert {c for _, c, _ in _filas_cliente(env)} == {None}


def test_avisa_si_el_sku_no_existe_en_el_sae_del_cliente(client, env, auth_as):
    auth_as(env["admin"]); h = _h(env); _activar(client, env)
    r = client.put(f"/api/v1/listas-precios/{env['lista']}/sku-cliente", headers=h,
                   json={"producto_id": env["prod"], "presentacion": "PIEZA", "sku": "ZANAHORIA-PZ-BJ"})
    assert r.status_code == 200
    avisos = r.json()["avisos"]
    assert len(avisos) == 1 and "BALLES" in avisos[0] and "SAE 02" in avisos[0]   # Jubran no está ligado a SAE aquí


def test_el_excel_de_la_lista_trae_y_sube_el_sku(client, env, auth_as):
    auth_as(env["admin"]); h = _h(env)
    base = f"/api/v1/listas-precios/{env['lista']}"
    # Sin la lista activada, el Excel no trae la columna.
    wb = load_workbook(io.BytesIO(client.get(f"{base}/export", headers=h).content))
    assert "SKU CLIENTE" not in [c.value for c in wb.active[1]]

    _activar(client, env)
    client.put(f"{base}/sku-cliente", headers=h,
               json={"producto_id": env["prod"], "presentacion": "KILO", "sku": "ZANA-FRUT-508"})
    wb = load_workbook(io.BytesIO(client.get(f"{base}/export", headers=h).content))
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    col = hdr.index("SKU CLIENTE")
    filas = {r[2]: r[col] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert filas == {"KILO": "ZANA-FRUT-508", "PIEZA": None}

    # Se sube: KILO cambia, PIEZA vacía no se toca.
    nuevo = Workbook(); w = nuevo.active
    w.append(hdr)
    for r in ws.iter_rows(min_row=2, values_only=True):
        r = list(r)
        if r[2] == "KILO":
            r[col] = "ZANA-NUEVO-1"
        w.append(r)
    buf = io.BytesIO(); nuevo.save(buf)
    r = client.post(f"{base}/importar", headers=h,
                    files={"archivo": ("lista.xlsx", buf.getvalue(),
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["skus_guardados"] == 1 and not res["errores"]
    assert len(res["avisos"]) == 1                      # no existe en el SAE 02 de Balles
    assert {c for _, c, _ in _filas_cliente(env)} == {"ZANA-NUEVO-1"}
