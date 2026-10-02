"""Envíos de cobranza: varias razones sociales en un estado de cuenta, la
tabla por proyecto/serie/sucursal/razón social, el previo, los archivos, la
prueba que solo llega a quien la pide y las razones sociales sin envío."""
import io
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal

import pytest

from app.core.db import SessionLocal
from app.models import Cliente, ClienteSucursal, Factura, Proyecto, Serie, Sucursal
from tests.test_cobranza_api import _h, auth, auth_atado, env  # noqa: F401
from tests.test_cobranza_auto import _config, correo  # noqa: F401

_BASE = "/api/v1/cobranza/automatica/grupos"


def _factura(env, cliente, *, total, dias_atras, folio, serie="F", notas=None, proyecto_id=None):
    db = SessionLocal()
    try:
        f = Factura(
            tenant_id=uuid.UUID(str(env["tenant_id"])), serie=serie, folio=folio,
            cliente_id=uuid.UUID(env[cliente]), metodo_pago="PPD", forma_pago="99",
            total=Decimal(str(total)), subtotal=Decimal(str(total)), estado="TIMBRADA",
            uuid=str(uuid.uuid4()), notas=notas,
            fecha=datetime.now(timezone.utc) - timedelta(days=dias_atras),
            saldo_insoluto=Decimal(str(total)),
            proyecto_id=uuid.UUID(proyecto_id) if proyecto_id else None,
        )
        db.add(f); db.commit()
        return str(f.id)
    finally:
        db.close()


def _plaza(env, nombre):
    db = SessionLocal()
    try:
        s = Sucursal(tenant_id=uuid.UUID(str(env["tenant_id"])), nombre=nombre)
        db.add(s); db.commit()
        return str(s.id)
    finally:
        db.close()


def _proyecto(env, nombre, *, cliente=None, series=(), sucursal=None):
    db = SessionLocal()
    try:
        p = Proyecto(tenant_id=uuid.UUID(str(env["tenant_id"])), codigo=nombre.replace(" ", "")[:20],
                     nombre=nombre, series=list(series),
                     cliente_id=uuid.UUID(env[cliente]) if cliente else None,
                     sucursal_id=uuid.UUID(sucursal) if sucursal else None)
        db.add(p); db.commit()
        return str(p.id)
    finally:
        db.close()


def _nombre_corto(env, cliente, nombre, **otros):
    db = SessionLocal()
    try:
        db.query(Cliente).filter(Cliente.id == uuid.UUID(env[cliente])).update({"nombre_corto": nombre, **otros})
        db.commit()
    finally:
        db.close()


@pytest.fixture
def ehmo(env):
    """Dos razones sociales como EHMO y SUREÑA: `cli` (30 días de crédito)
    con dos proyectos en plazas distintas y una serie sin proyecto; `otro`
    (15 días) con un proyecto sin sucursal. Las razones sociales van en orden
    alfabético de su razón social (Cliente Cobranza < Sureña del Sur)."""
    _nombre_corto(env, "cli", "EHMO")
    _nombre_corto(env, "otro", "SUR", legal_name="Sureña del Sur")
    hgo = _plaza(env, "Hidalgo")
    chis = _plaza(env, "Chiapas")
    p = {
        "hos": _proyecto(env, "HOSPITALES HIDALGO", cliente="cli", series=["ZHOS"], sucursal=hgo),
        "tg": _proyecto(env, "HOSPITALES TUXTLA", cliente="cli", series=["ZTG"], sucursal=chis),
        "com": _proyecto(env, "COMEDORES", cliente="otro", series=["ZSUR"]),
        "hgo": hgo, "chis": chis,
    }
    _factura(env, "cli", total=1000, dias_atras=10, folio=1, serie="ZHOS")    # por vencer (30 d)
    _factura(env, "cli", total=2000, dias_atras=40, folio=2, serie="ZHOS")    # vencida
    _factura(env, "cli", total=500, dias_atras=40, folio=3, serie="ZTG")      # vencida
    _factura(env, "cli", total=300, dias_atras=5, folio=4, serie="ZSUELTA")   # sin proyecto
    _factura(env, "otro", total=700, dias_atras=20, folio=5, serie="ZSUR")    # vencida (15 d)
    return p


def _grupo(env, **kw):
    return {"nombre": "EHMO", "agrupar_por": "PROYECTO", "correos": ["pagos@ehmo.example.com"],
            "alcance": [{"cliente_id": env["cli"]}, {"cliente_id": env["otro"]}], **kw}


def _previo(client, env, cuerpo):
    r = client.post(f"{_BASE}/previo", json=cuerpo, headers=_h(env))
    assert r.status_code == 200, r.text
    return r.json()


def test_por_proyecto_junta_las_razones_sociales(client, env, auth, ehmo):
    _config(client, env, cc_siempre=["cobranza@negocio.example.com"])
    d = _previo(client, env, _grupo(env))
    r = d["resumen"]
    assert float(r["saldo_total"]) == 4500.0
    # Cada factura vence con los días de SU razón social: la de SUR (15 días,
    # hace 20) ya venció; la de EHMO de hace 10 días (30) no.
    assert float(r["vencido_total"]) == 3200.0
    filas = [(f["clientes"], f["series"], f["proyectos"], float(f["saldo"])) for f in r["filas"]]
    assert filas == [
        (["EHMO"], ["ZHOS"], ["HOSPITALES HIDALGO"], 3000.0),
        (["EHMO"], ["ZTG"], ["HOSPITALES TUXTLA"], 500.0),
        (["EHMO"], ["ZSUELTA"], [], 300.0),          # sin proyecto, al final de su razón social
        (["SUR"], ["ZSUR"], ["COMEDORES"], 700.0),
    ]
    assert d["para"] == ["pagos@ehmo.example.com"]
    assert d["cc"] == ["cobranza@negocio.example.com"]
    assert "HOSPITALES TUXTLA" in d["html"] and "Total EHMO" in d["html"]
    assert d["hojas"] == ["Resumen", "HOSPITALES HIDALGO", "HOSPITALES TUXTLA", "EHMO ZSUELTA", "COMEDORES"]
    assert d["avisos"] == []


def test_otras_agrupaciones(client, env, auth, ehmo):
    _config(client, env)
    por = lambda modo: _previo(client, env, _grupo(env, agrupar_por=modo))["resumen"]["filas"]  # noqa: E731

    serie = por("SERIE")
    assert [f["series"] for f in serie] == [["ZHOS"], ["ZSUELTA"], ["ZTG"], ["ZSUR"]]

    cliente = por("CLIENTE")
    assert [(f["clientes"], float(f["saldo"])) for f in cliente] == [(["EHMO"], 3800.0), (["SUR"], 700.0)]

    d = _previo(client, env, _grupo(env, agrupar_por="SUCURSAL"))
    suc = d["resumen"]["filas"]
    # ZSUELTA no tiene proyecto ni vínculo: también cae en «Sin sucursal».
    assert [(f["hoja"], float(f["saldo"])) for f in suc] == [
        ("Chiapas", 500.0), ("Hidalgo", 3000.0), ("Sin sucursal", 1000.0)]
    # El proyecto sin sucursal se avisa por nombre.
    assert any("COMEDORES" in a for a in d["avisos"])


def test_sucursal_de_lo_que_no_tiene_proyecto_sale_del_vinculo(client, env, auth, ehmo):
    _config(client, env)
    db = SessionLocal()
    try:
        tid = uuid.UUID(str(env["tenant_id"]))
        s = Serie(tenant_id=tid, codigo="ZSUELTA", tipo_documento="FACTURA")
        db.add(s); db.flush()
        db.add(ClienteSucursal(tenant_id=tid, cliente_id=uuid.UUID(env["cli"]),
                               sucursal_id=uuid.UUID(ehmo["hgo"]), serie_factura_id=s.id))
        db.commit()
    finally:
        db.close()
    filas = _previo(client, env, _grupo(env, agrupar_por="SUCURSAL"))["resumen"]["filas"]
    assert [(f["hoja"], float(f["saldo"])) for f in filas] == [
        ("Chiapas", 500.0), ("Hidalgo", 3300.0), ("Sin sucursal", 700.0)]


def test_alcance_parcial_solo_lo_marcado(client, env, auth, ehmo):
    _config(client, env)
    cuerpo = _grupo(env, alcance=[{"cliente_id": env["cli"], "completo": False,
                                   "proyectos": [ehmo["tg"]], "series": ["zsuelta"]}])
    filas = _previo(client, env, cuerpo)["resumen"]["filas"]
    assert [(f["series"], float(f["saldo"])) for f in filas] == [(["ZTG"], 500.0), (["ZSUELTA"], 300.0)]

    vacio = _grupo(env, alcance=[{"cliente_id": env["cli"], "completo": False}])
    r = client.post(f"{_BASE}/previo", json=vacio, headers=_h(env))
    assert r.status_code == 422 and "EHMO" in r.json()["detail"]


def test_antiguedad_y_solo_vencidas(client, env, auth, ehmo):
    _config(client, env)
    d = _previo(client, env, _grupo(env, mostrar_antiguedad=True, incluir_por_vencer=False))
    assert float(d["resumen"]["saldo_total"]) == 3200.0      # sin las por vencer
    assert "Más de 90" in d["html"] and "1 a 30" in d["html"]
    assert any("Solo van las facturas vencidas" in a for a in d["avisos"])


def test_crud_y_tambien_en(client, env, auth, ehmo):
    _config(client, env)
    r = client.post(_BASE, json=_grupo(env), headers=_h(env))
    assert r.status_code == 201, r.text
    g = r.json()
    assert float(g["saldo"]) == 4500.0 and g["tambien_en"] == []
    assert [a["cliente"] for a in g["alcance"]] == ["EHMO", "SUR"]

    # Mismo nombre (sin importar mayúsculas): 409.
    r = client.post(_BASE, json=_grupo(env, nombre="ehmo"), headers=_h(env))
    assert r.status_code == 409

    # Un segundo grupo con solo Tuxtla choca con el primero (EHMO completo).
    tux = _grupo(env, nombre="EHMO Tuxtla",
                 alcance=[{"cliente_id": env["cli"], "completo": False, "proyectos": [ehmo["tg"]]}])
    r = client.post(_BASE, json=tux, headers=_h(env))
    assert r.status_code == 201, r.text
    lista = {x["nombre"]: x for x in client.get(_BASE, headers=_h(env)).json()["envios"]}
    assert [t["nombre"] for t in lista["EHMO"]["tambien_en"]] == ["EHMO Tuxtla"]
    assert float(lista["EHMO Tuxtla"]["saldo"]) == 500.0
    # El previo también lo avisa.
    d = _previo(client, env, {**_grupo(env), "id": g["id"]})
    assert any("EHMO Tuxtla" in a for a in d["avisos"])

    # Cambiar: solo SUR, por razón social, automático los viernes, con nota.
    r = client.put(f"{_BASE}/{g['id']}", headers=_h(env), json=_grupo(
        env, agrupar_por="CLIENTE", modo="AUTOMATICO", dia_semana=4, hora=10, nota="convenio",
        alcance=[{"cliente_id": env["otro"]}]))
    assert r.status_code == 200, r.text
    out = r.json()
    assert (out["agrupar_por"], out["modo"], out["dia_semana"], out["hora"], out["nota"]) == (
        "CLIENTE", "AUTOMATICO", 4, 10, "convenio")
    assert float(out["saldo"]) == 700.0 and out["tambien_en"] == []

    assert client.delete(f"{_BASE}/{g['id']}", headers=_h(env)).status_code == 200
    assert [x["nombre"] for x in client.get(_BASE, headers=_h(env)).json()["envios"]] == ["EHMO Tuxtla"]


def test_validaciones(client, env, auth, ehmo):
    r = client.post(_BASE, json=_grupo(env, correos=["no-es-correo"]), headers=_h(env))
    assert r.status_code == 422
    r = client.post(_BASE, json=_grupo(env, alcance=[{"cliente_id": env["cli"]}, {"cliente_id": env["cli"]}]),
                    headers=_h(env))
    assert r.status_code == 422
    r = client.post(_BASE, json=_grupo(env, alcance=[{"cliente_id": str(uuid.uuid4())}]), headers=_h(env))
    assert r.status_code == 404


def test_candado_por_cliente(client, env, auth_atado, ehmo):
    """El usuario amarrado a `cli` no arma grupos con razones que no ve."""
    r = client.post(_BASE, json=_grupo(env), headers=_h(env))
    assert r.status_code == 404
    r = client.post(_BASE, json=_grupo(env, alcance=[{"cliente_id": env["cli"]}]), headers=_h(env))
    assert r.status_code == 201, r.text
    ops = client.get(f"{_BASE}/opciones", headers=_h(env)).json()
    assert [o["cliente_id"] for o in ops] == [env["cli"]]


def test_opciones_del_arbol(client, env, auth, ehmo):
    ops = {o["nombre"]: o for o in client.get(f"{_BASE}/opciones", headers=_h(env)).json()}
    nodos = [(n["proyecto"], n["serie"], n["sucursal"], float(n["saldo"])) for n in ops["EHMO"]["nodos"]]
    assert nodos == [("HOSPITALES HIDALGO", None, "Hidalgo", 3000.0),
                     ("HOSPITALES TUXTLA", None, "Chiapas", 500.0),
                     (None, "ZSUELTA", None, 300.0)]
    assert float(ops["SUR"]["saldo"]) == 700.0


def test_excel_y_pdf(client, env, auth, ehmo):
    from openpyxl import load_workbook

    _config(client, env)
    r = client.post(f"{_BASE}/previo/xlsx", json=_grupo(env), headers=_h(env))
    assert r.status_code == 200, r.text
    assert 'filename="estado-cuenta-EHMO-' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Resumen", "HOSPITALES HIDALGO", "HOSPITALES TUXTLA", "EHMO ZSUELTA", "COMEDORES"]
    hoja = wb["HOSPITALES HIDALGO"]
    folios = [row[5] for row in hoja.iter_rows(values_only=True) if row and row[5] and str(row[5]).startswith("ZHOS")]
    assert folios == ["ZHOS 2", "ZHOS 1"]                    # por fecha, la más vieja primero
    resumen = [row for row in wb["Resumen"].iter_rows(values_only=True) if row and row[0] == "SUR"]
    assert resumen[0][3] == 700.0

    r = client.post(f"{_BASE}/previo/pdf", json=_grupo(env), headers=_h(env))
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_prueba_solo_a_quien_la_pide(client, env, auth, ehmo, correo):
    _config(client, env)
    r = client.post(f"{_BASE}/prueba", json=_grupo(env), headers=_h(env))
    assert r.status_code == 200, r.text
    assert correo[0]["to"] == [env["email"]] and correo[0]["cc"] is None
    assert correo[0]["subject"].startswith("[Prueba] Estado de cuenta EHMO-SUR al ")
    assert sorted(a.rsplit(".", 1)[1] for a in correo[0]["adjuntos"]) == ["pdf", "xlsx"]


def test_nombre_de_la_tabla(client, env, auth, ehmo):
    """El envío «GRUPO OPERADOR DE ALIMENTOS EHMO» juntaba EHMO, SUREÑA y
    MAFAN y la tabla cerraba con «Total GRUPO OPERADOR…». Lo que lee el
    cliente dice el nombre de la tabla; vacío, los nombres cortos con guion."""
    from openpyxl import load_workbook

    _config(client, env)
    cuerpo = _grupo(env, nombre="GRUPO OPERADOR DE ALIMENTOS EHMO")
    d = _previo(client, env, cuerpo)
    assert d["resumen"]["titulo"] == "EHMO-SUR"
    assert "Total EHMO-SUR" in d["html"] and "Total GRUPO" not in d["html"]
    assert d["asunto"].startswith("Estado de cuenta EHMO-SUR al ")
    assert d["adjuntos"][0]["nombre"].startswith("estado-cuenta-EHMO-SUR-")

    # Capturado, manda en el correo y en los archivos.
    cuerpo["titulo_tabla"] = "  Saldos   EHMO "
    d = _previo(client, env, cuerpo)
    assert d["resumen"]["titulo"] == "Saldos EHMO" and "Total Saldos EHMO" in d["html"]
    r = client.post(f"{_BASE}/previo/xlsx", json=cuerpo, headers=_h(env))
    total = [row[0] for row in load_workbook(io.BytesIO(r.content))["Resumen"].iter_rows(values_only=True)
             if row and str(row[0] or "").startswith("TOTAL")]
    assert total == ["TOTAL SALDOS EHMO"]

    # Se guarda y vuelve en la lista; borrado, regresa al prellenado.
    r = client.post(_BASE, json=cuerpo, headers=_h(env))
    assert r.status_code == 201, r.text
    gid = r.json()["id"]
    envio = next(e for e in client.get(_BASE, headers=_h(env)).json()["envios"] if e["id"] == gid)
    assert envio["titulo_tabla"] == "Saldos EHMO"
    r = client.put(f"{_BASE}/{gid}", json={**cuerpo, "titulo_tabla": "  "}, headers=_h(env))
    assert r.status_code == 200, r.text
    envio = next(e for e in client.get(_BASE, headers=_h(env)).json()["envios"] if e["id"] == gid)
    assert envio["titulo_tabla"] is None


def test_titulo_sugerido():
    from app.services.cobranza_grupos import titulo_sugerido

    assert titulo_sugerido("EHMO", ["SUREÑA", "EHMO", "MAFAN"]) == "EHMO-MAFAN-SUREÑA"
    assert titulo_sugerido("EHMO", ["EHMO", " EHMO "]) == "EHMO"
    # Diez razones sociales ya no se leen como nombre: toma el del envío.
    largas = [f"RAZON SOCIAL NUMERO {i}" for i in range(10)]
    assert titulo_sugerido("REPORTE CRISTIAN", largas) == "REPORTE CRISTIAN"


def test_razones_sociales_sin_envio(client, env, auth, ehmo):
    _config(client, env)
    sin = {x["nombre"]: x for x in client.get(_BASE, headers=_h(env)).json()["sin_envio"]}
    assert float(sin["EHMO"]["saldo"]) == 3800.0 and float(sin["SUR"]["saldo"]) == 700.0
    # Un envío con solo Tuxtla deja a EHMO cubierta en parte.
    client.post(_BASE, headers=_h(env), json=_grupo(
        env, alcance=[{"cliente_id": env["cli"], "completo": False, "proyectos": [ehmo["tg"]]}]))
    sin = {x["nombre"]: x for x in client.get(_BASE, headers=_h(env)).json()["sin_envio"]}
    assert float(sin["EHMO"]["saldo"]) == 3300.0 and sin["EHMO"]["parcial"] is True
    assert "SUR" in sin


def test_nombre_corto_en_la_ficha(client, env, auth):
    r = client.patch(f"/api/v1/clientes/{env['cli']}", json={"nombre_corto": "  EHMO   Sur "}, headers=_h(env))
    assert r.status_code == 200, r.text
    assert r.json()["nombre_corto"] == "EHMO Sur"
    r = client.patch(f"/api/v1/clientes/{env['cli']}", json={"nombre_corto": "  "}, headers=_h(env))
    assert r.json()["nombre_corto"] is None
