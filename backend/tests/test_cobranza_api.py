"""Cobranza F1 — estado de cuenta del cliente con antigüedad de saldos.

Verifica: saldo_insoluto se inicializa al timbrar (PPD=total, PUE=0) y el
estado de cuenta agrupa por fecha de vencimiento (fecha + días de crédito) en
cubetas de 30 días con "por vencer", estilo SAE.
"""
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal

import pytest
from sqlalchemy import text

from app.core.auth import Principal, get_principal
from app.core.db import SessionLocal
from app.core.rbac import invalidate_auth_cache
from app.main import app
from app.models import Cliente, Factura, Membership, Role, Tenant, User

_PURGE = ("recibo_pago_facturas", "recibos_pago", "timbrado_intentos", "lineas_factura", "facturas", "clientes")


@pytest.fixture
def env(db_engine):
    suffix = uuid.uuid4().hex[:8]
    db = SessionLocal()
    created = {"m": [], "u": [], "t": []}
    try:
        tenant = Tenant(slug=f"cob-{suffix}", legal_name="COB SA", rfc=f"C{suffix.upper()}B"[:13],
                        regimen_fiscal_sat="601", domicilio_fiscal_cp="44100",
                        tier="PRINCIPAL", status="ACTIVE")
        db.add(tenant); db.flush(); created["t"].append(tenant.id)
        role = db.query(Role).filter(Role.nombre == "ADMIN", Role.es_preset.is_(True)).one()
        cli = Cliente(tenant_id=tenant.id, codigo="COBC", legal_name="Cliente Cobranza",
                      rfc="XAXX010101000", regimen_fiscal="601", uso_cfdi_default="G03",
                      dias_credito=30, limite_credito=Decimal("100000"))
        # El segundo cliente existe solo para probar el candado: es "el de otro".
        otro = Cliente(tenant_id=tenant.id, codigo="COBO", legal_name="Cliente Ajeno",
                       rfc="XEXX010101000", regimen_fiscal="601", uso_cfdi_default="G03",
                       dias_credito=15)
        db.add_all([cli, otro]); db.flush()

        def _user(label, scope=None):
            sub = f"sub-{label}-{suffix}"
            u = User(email=f"{label}-{suffix}@t.test", auth_user_id=sub, full_name=label)
            db.add(u); db.flush(); created["u"].append(u.id)
            m = Membership(tenant_id=tenant.id, user_id=u.id, role_id=role.id,
                           cliente_scope=scope)
            db.add(m); db.flush(); created["m"].append(m.id)
            return {"sub": sub, "email": u.email, "tenant_id": tenant.id}

        admin = _user("cob")
        # Mismos permisos, pero amarrado a un cliente: el candado es de la
        # membresía, no del rol.
        atado = _user("cobscope", scope=[cli.id])
        db.commit()
        invalidate_auth_cache()
        yield {**admin, "cli": str(cli.id), "otro": str(otro.id), "atado": atado}
    finally:
        for tb in _PURGE:
            for tid in created["t"]:
                db.execute(text(f"DELETE FROM {tb} WHERE tenant_id = :t"), {"t": tid})
        for mid in created["m"]:
            db.query(Membership).filter(Membership.id == mid).delete()
        for uid in created["u"]:
            db.query(User).filter(User.id == uid).delete()
        for tid in created["t"]:
            db.query(Tenant).filter(Tenant.id == tid).delete()
        db.commit(); db.close()
        invalidate_auth_cache()


def _como(usuario):
    app.dependency_overrides[get_principal] = lambda: Principal(
        auth_user_id=usuario["sub"], email=usuario["email"], role="authenticated",
        claims={"sub": usuario["sub"]})


@pytest.fixture
def auth(env):
    _como(env)
    yield
    app.dependency_overrides.pop(get_principal, None)


@pytest.fixture
def auth_atado(env):
    """Sesión del usuario amarrado al cliente `cli` (candado por cliente)."""
    _como(env["atado"])
    yield
    app.dependency_overrides.pop(get_principal, None)


def _h(env):
    return {"X-Tenant-Id": str(env["tenant_id"])}


def _factura_ppd_timbrada(env, *, total, dias_atras, metodo="PPD", folio, serie="F", notas=None):
    """Inserta una factura TIMBRADA directamente (sin PAC) con fecha dada."""
    db = SessionLocal()
    try:
        f = Factura(
            tenant_id=uuid.UUID(str(env["tenant_id"])), serie=serie, folio=folio,
            cliente_id=uuid.UUID(env["cli"]), metodo_pago=metodo, forma_pago="99",
            total=Decimal(str(total)), subtotal=Decimal(str(total)),
            estado="TIMBRADA", uuid=str(uuid.uuid4()), notas=notas,
            fecha=datetime.now(timezone.utc) - timedelta(days=dias_atras),
            saldo_insoluto=Decimal(str(total)) if metodo == "PPD" else Decimal("0"),
        )
        db.add(f); db.commit()
        return str(f.id)
    finally:
        db.close()


def test_estado_cuenta_antiguedad_por_vencimiento(client, env, auth):
    # Cliente con 30 días de crédito. Facturas:
    #  - $1000 hace 10 días  → vence en +20 → POR VENCER
    #  - $2000 hace 40 días  → venció hace 10 → 1-30
    #  - $3000 hace 100 días → venció hace 70 → 61-90
    #  - $500 PUE hace 5 días → saldo 0, NO aparece
    _factura_ppd_timbrada(env, total=1000, dias_atras=10, folio=1)
    _factura_ppd_timbrada(env, total=2000, dias_atras=40, folio=2)
    _factura_ppd_timbrada(env, total=3000, dias_atras=100, folio=3)
    _factura_ppd_timbrada(env, total=500, dias_atras=5, metodo="PUE", folio=4)

    r = client.get(f"/api/v1/cobranza/estado-cuenta/{env['cli']}", headers=_h(env))
    assert r.status_code == 200, r.text
    d = r.json()
    assert float(d["saldo_total"]) == 6000.0            # PUE no cuenta
    assert len(d["facturas"]) == 3
    a = d["antiguedad"]
    assert float(a["por_vencer"]) == 1000.0
    assert float(a["d1_30"]) == 2000.0
    assert float(a["d61_90"]) == 3000.0
    assert float(a["d31_60"]) == 0.0 and float(a["d90_mas"]) == 0.0
    assert d["dias_credito"] == 30


def test_estado_cuenta_vacio(client, env, auth):
    r = client.get(f"/api/v1/cobranza/estado-cuenta/{env['cli']}", headers=_h(env))
    assert r.status_code == 200, r.text
    assert float(r.json()["saldo_total"]) == 0.0
    assert r.json()["facturas"] == []


def test_extraer_semana():
    """La semana de entrega: dicha con letras en la observación de SAE o
    embebida en el folio interno (ahí el número ES la semana)."""
    from app.services.espejo_cruce import extraer_semana

    assert extraer_semana("SEMANA 26 HUASTECA LUNES 29 /JUNIO/2026") == 26
    assert extraer_semana("SEMANA 33 SECRETARIO NERI REQ 20/08/2026 SN-33NER-JUE") == 33
    assert extraer_semana("entrega HO-34VIL-MIE pactada") == 34
    assert extraer_semana("CEN-35HUA-EMB") == 35          # proyecto de 3 letras
    assert extraer_semana(None, "HO-34VIL-MIE") == 34     # fallback al su_pedido
    assert extraer_semana("OC 0000024736 ENTREGA CEDIS") is None
    assert extraer_semana(None, "") is None


def test_estado_cuenta_filtra_por_serie_y_deriva_semana(client, env, auth):
    """`serie` acota facturas y totales; `series` SIEMPRE trae el resumen
    completo (para el filtro de la pantalla); la semana sale de las notas."""
    _factura_ppd_timbrada(env, total=1000, dias_atras=5, folio=1, serie="ZEH",
                          notas="SEMANA 26 HUASTECA LUNES 29 /JUNIO/2026")
    _factura_ppd_timbrada(env, total=2000, dias_atras=5, folio=2, serie="FEH",
                          notas="ENTREGA HO-34VIL-MIE")

    r = client.get(f"/api/v1/cobranza/estado-cuenta/{env['cli']}?serie=ZEH", headers=_h(env))
    assert r.status_code == 200, r.text
    d = r.json()
    assert [f["folio"] for f in d["facturas"]] == [1]
    assert float(d["saldo_total"]) == 1000.0
    assert d["facturas"][0]["semana"] == 26
    assert {s["serie"]: s["facturas"] for s in d["series"]} == {"ZEH": 1, "FEH": 1}

    todas = client.get(f"/api/v1/cobranza/estado-cuenta/{env['cli']}", headers=_h(env)).json()
    assert float(todas["saldo_total"]) == 3000.0
    assert {f["folio"]: f["semana"] for f in todas["facturas"]} == {1: 26, 2: 34}


def test_estado_cuenta_xlsx(client, env, auth):
    """El Excel estilo SAE: encabezado con crédito, la tabla con SEM/FACTURA/
    SALDOS y el renglón TOTAL con fórmulas =SUM."""
    import io

    from openpyxl import load_workbook

    _factura_ppd_timbrada(env, total=1500, dias_atras=40, folio=7, serie="ZEH",
                          notas="SEMANA 26 HUASTECA LUNES 29 /JUNIO/2026")
    _factura_ppd_timbrada(env, total=500, dias_atras=1, folio=8, serie="ZEH")

    r = client.get(f"/api/v1/cobranza/estado-cuenta/{env['cli']}/xlsx?serie=ZEH", headers=_h(env))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "estado-cuenta-ZEH-" in r.headers["content-disposition"]

    ws = load_workbook(io.BytesIO(r.content)).active
    celdas = {c.value for row in ws.iter_rows() for c in row if c.value is not None}
    assert "Cliente Cobranza" in celdas          # membrete
    assert "Días de crédito:" in celdas
    assert {"SEM", "CLIENTE", "PROYECTO", "FACTURA", "FECHA APLIC",
            "SUBTOTAL", "ABONOS", "SALDOS", "ESTATUS"} <= celdas
    assert {"ZEH 7", "ZEH 8", "SEM 26", "VENCIDO"} <= celdas  # 40 días > 30 de crédito
    # El renglón TOTAL suma con fórmula, no con un número pegado.
    ultima = [c.value for c in ws[ws.max_row]]
    assert any(isinstance(v, str) and v.startswith("=SUM(") for v in ultima)


# ── Recibos de Pago (REP) F2 ─────────────────────────────────────────────────
import app.api.v1.cobranza as cobranza_mod


class _FakePAC:
    configured = True
    env_label = "sandbox"

    @classmethod
    def from_settings(cls, settings):
        return cls()

    def create_cfdi_pago(self, payload):
        return {"Id": "REP-FAKE", "Complement": {"TaxStamp": {"Uuid": str(uuid.uuid4())}}}

    def cancel_cfdi(self, cfdi_id, motive, uuid_replacement=None):
        return {"Status": "canceled"}

    def download_xml(self, cfdi_id):
        return b"<xml/>"

    def buscar_cfdi(self, order_number, **kw):
        return True, None


@pytest.fixture
def fake_pac(monkeypatch):
    monkeypatch.setattr(cobranza_mod, "FacturamaClient", _FakePAC)


def _saldo_factura(fid):
    db = SessionLocal()
    try:
        return Decimal(db.query(Factura).filter(Factura.id == uuid.UUID(fid)).one().saldo_insoluto)
    finally:
        db.close()


def test_rep_registrar_timbrar_descuenta_saldo(client, env, auth, fake_pac):
    # Factura PPD de $1160 con saldo full.
    fid = _factura_ppd_timbrada(env, total=1160, dias_atras=40, folio=10)
    h = _h(env)
    hoy = datetime.now(timezone.utc).isoformat()
    # Pago parcial de $500 → abono, saldo 660.
    r = client.post("/api/v1/cobranza/recibos-pago", headers=h, json={
        "cliente_id": env["cli"], "fecha_pago": hoy, "forma_pago": "03", "monto": "500",
        "facturas": [{"factura_id": fid, "importe": "500"}]})
    assert r.status_code == 201, r.text
    rec = r.json()
    assert rec["estado"] == "BORRADOR"
    assert rec["facturas"][0]["num_parcialidad"] == 1
    assert float(rec["facturas"][0]["saldo_anterior"]) == 1160.0
    assert float(rec["facturas"][0]["saldo_insoluto"]) == 660.0
    assert _saldo_factura(fid) == Decimal("1160")     # aún no timbrado, no descuenta

    t = client.post(f"/api/v1/cobranza/recibos-pago/{rec['id']}/timbrar", headers=h)
    assert t.status_code == 200, t.text
    assert t.json()["estado"] == "TIMBRADO" and t.json()["uuid"]
    assert _saldo_factura(fid) == Decimal("660")      # timbrado → descuenta

    # Segundo abono de $660 → parcialidad 2, salda la factura.
    r2 = client.post("/api/v1/cobranza/recibos-pago", headers=h, json={
        "cliente_id": env["cli"], "fecha_pago": hoy, "forma_pago": "03", "monto": "660",
        "facturas": [{"factura_id": fid, "importe": "660"}]})
    assert r2.json()["facturas"][0]["num_parcialidad"] == 2
    t2 = client.post(f"/api/v1/cobranza/recibos-pago/{r2.json()['id']}/timbrar", headers=h)
    assert t2.status_code == 200, t2.text
    assert _saldo_factura(fid) == Decimal("0")        # saldada
    # Ya no aparece en el estado de cuenta.
    ec = client.get(f"/api/v1/cobranza/estado-cuenta/{env['cli']}", headers=h).json()
    assert all(d["factura_id"] != fid for d in ec["facturas"])


def test_rep_valida_monto_y_saldo(client, env, auth):
    fid = _factura_ppd_timbrada(env, total=1000, dias_atras=10, folio=11)
    h = _h(env); hoy = datetime.now(timezone.utc).isoformat()
    # suma no cuadra con monto → 422
    r = client.post("/api/v1/cobranza/recibos-pago", headers=h, json={
        "cliente_id": env["cli"], "fecha_pago": hoy, "forma_pago": "03", "monto": "500",
        "facturas": [{"factura_id": fid, "importe": "300"}]})
    assert r.status_code == 422 and "cuadra" in r.json()["detail"]
    # importe excede saldo → 422
    r2 = client.post("/api/v1/cobranza/recibos-pago", headers=h, json={
        "cliente_id": env["cli"], "fecha_pago": hoy, "forma_pago": "03", "monto": "2000",
        "facturas": [{"factura_id": fid, "importe": "2000"}]})
    assert r2.status_code == 422 and "excede el saldo" in r2.json()["detail"]


def test_rep_solo_ppd_timbrada(client, env, auth):
    # PUE timbrada → no aplica REP.
    fid = _factura_ppd_timbrada(env, total=500, dias_atras=5, metodo="PUE", folio=12)
    h = _h(env); hoy = datetime.now(timezone.utc).isoformat()
    r = client.post("/api/v1/cobranza/recibos-pago", headers=h, json={
        "cliente_id": env["cli"], "fecha_pago": hoy, "forma_pago": "03", "monto": "500",
        "facturas": [{"factura_id": fid, "importe": "500"}]})
    assert r.status_code == 422 and "PPD timbrada" in r.json()["detail"]


def test_rep_gate_multiemisor_sin_csd_es_422(client, env, auth, fake_pac, monkeypatch):
    """Con multiemisor prendido, un tenant que no está listo (sin CSD propio /
    datos fiscales) recibe el mismo 422 accionable que timbrar factura — no el
    error crudo del PAC (502)."""
    import app.core.config as cfg
    monkeypatch.setattr(cfg.settings, "FACTURAMA_MULTIEMISOR", True)
    fid = _factura_ppd_timbrada(env, total=100, dias_atras=1, folio=77)
    h = _h(env); hoy = datetime.now(timezone.utc).isoformat()
    r = client.post("/api/v1/cobranza/recibos-pago", headers=h, json={
        "cliente_id": env["cli"], "fecha_pago": hoy, "forma_pago": "03", "monto": "100",
        "facturas": [{"factura_id": fid, "importe": "100"}]})
    assert r.status_code == 201, r.text
    t = client.post(f"/api/v1/cobranza/recibos-pago/{r.json()['id']}/timbrar", headers=h)
    assert t.status_code == 422, t.text
    assert "no está lista para facturar" in t.json()["detail"]


def test_rep_gate_multiemisor_listo_timbra(client, env, auth, monkeypatch):
    """Con datos fiscales completos y el CSD del RFC del tenant en Facturama,
    el gate deja pasar y el REP se timbra igual que antes."""
    import app.core.config as cfg
    monkeypatch.setattr(cfg.settings, "FACTURAMA_MULTIEMISOR", True)

    rfc = "COB800101AA1"           # formato SAT válido (el del fixture no lo es)
    db = SessionLocal()
    try:
        db.query(Tenant).filter(Tenant.id == env["tenant_id"]).update({"rfc": rfc})
        db.commit()
    finally:
        db.close()

    class _PACConCSD(_FakePAC):
        def listar_csds(self):
            return [{"Rfc": rfc}]

    monkeypatch.setattr(cobranza_mod, "FacturamaClient", _PACConCSD)

    fid = _factura_ppd_timbrada(env, total=200, dias_atras=1, folio=78)
    h = _h(env); hoy = datetime.now(timezone.utc).isoformat()
    r = client.post("/api/v1/cobranza/recibos-pago", headers=h, json={
        "cliente_id": env["cli"], "fecha_pago": hoy, "forma_pago": "03", "monto": "200",
        "facturas": [{"factura_id": fid, "importe": "200"}]})
    assert r.status_code == 201, r.text
    t = client.post(f"/api/v1/cobranza/recibos-pago/{r.json()['id']}/timbrar", headers=h)
    assert t.status_code == 200, t.text
    assert t.json()["estado"] == "TIMBRADO"


# ── F3: cancelación del REP + candado en cancelar_factura ────────────────────
def _timbrar_rep(client, env, h, fid, *, monto):
    hoy = datetime.now(timezone.utc).isoformat()
    r = client.post("/api/v1/cobranza/recibos-pago", headers=h, json={
        "cliente_id": env["cli"], "fecha_pago": hoy, "forma_pago": "03", "monto": monto,
        "facturas": [{"factura_id": fid, "importe": monto}]})
    assert r.status_code == 201, r.text
    rid = r.json()["id"]
    t = client.post(f"/api/v1/cobranza/recibos-pago/{rid}/timbrar", headers=h)
    assert t.status_code == 200, t.text
    return rid


def test_rep_cancelar_revierte_saldo(client, env, auth, fake_pac):
    # Factura $1160, abono $500 → saldo 660; al cancelar el REP el saldo regresa a 1160.
    fid = _factura_ppd_timbrada(env, total=1160, dias_atras=40, folio=20)
    h = _h(env)
    rid = _timbrar_rep(client, env, h, fid, monto="500")
    assert _saldo_factura(fid) == Decimal("660")

    c = client.post(f"/api/v1/cobranza/recibos-pago/{rid}/cancelar", headers=h, json={"motivo": "02"})
    assert c.status_code == 200, c.text
    assert c.json()["estado"] == "CANCELADO"
    assert _saldo_factura(fid) == Decimal("1160")     # abono revertido

    # No se cancela dos veces.
    c2 = client.post(f"/api/v1/cobranza/recibos-pago/{rid}/cancelar", headers=h, json={"motivo": "02"})
    assert c2.status_code == 409


def test_candado_factura_ppd_con_rep(client, env, auth, fake_pac, monkeypatch):
    # Con un REP timbrado, la factura no se cancela; tras cancelar el REP, sí.
    import app.core.config as cfg
    monkeypatch.setattr(cfg.settings, "FACTURAMA_FAKE_CANCEL", True)
    fid = _factura_ppd_timbrada(env, total=1000, dias_atras=10, folio=21)
    h = _h(env)
    rid = _timbrar_rep(client, env, h, fid, monto="1000")
    assert _saldo_factura(fid) == Decimal("0")

    blocked = client.post(f"/api/v1/facturas/{fid}/cancelar", headers=h, json={})
    assert blocked.status_code == 409 and "REP" in blocked.json()["detail"]

    # Cancelado el REP, la factura ya se puede cancelar.
    client.post(f"/api/v1/cobranza/recibos-pago/{rid}/cancelar", headers=h, json={"motivo": "02"})
    ok = client.post(f"/api/v1/facturas/{fid}/cancelar", headers=h, json={})
    assert ok.status_code == 200, ok.text
    assert ok.json()["estado"] == "CANCELADA"


# ── Estado de cuenta en PDF y por correo (corte de la migración) ─────────────
def _configura_correo(env):
    """Deja SMTP utilizable en el tenant: `configured()` exige host+user+pass."""
    db = SessionLocal()
    try:
        t = db.query(Tenant).filter(Tenant.id == env["tenant_id"]).one()
        t.config = {**(t.config or {}), "email": {
            "host": "smtp.test", "port": 587, "username": "envios@t.test",
            "password": "secreto", "from_name": "COB SA"}}
        db.commit()
    finally:
        db.close()


def test_estado_cuenta_pdf(client, env, auth):
    _factura_ppd_timbrada(env, total=1000, dias_atras=10, folio=30)
    _factura_ppd_timbrada(env, total=2000, dias_atras=100, folio=31)
    r = client.get(f"/api/v1/cobranza/estado-cuenta/{env['cli']}/pdf", headers=_h(env))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert r.content[:4] == b"%PDF"

    # Sin cargos también sale (el cliente al corriente recibe su hoja en cero).
    vacio = client.get(f"/api/v1/cobranza/estado-cuenta/{env['otro']}/pdf", headers=_h(env))
    assert vacio.status_code == 200 and vacio.content[:4] == b"%PDF"


def test_enviar_estado_cuenta_adjunta_el_pdf(client, env, auth, monkeypatch):
    from app.services import email as email_service

    _configura_correo(env)
    _factura_ppd_timbrada(env, total=1500, dias_atras=45, folio=32)
    enviado = {}

    def _fake_send(cfg, to, subject, html, attachments=None, reply_to=None):
        enviado.update({"to": to, "subject": subject, "html": html, "adj": attachments})

    monkeypatch.setattr(email_service, "send_email", _fake_send)
    r = client.post(f"/api/v1/cobranza/estado-cuenta/{env['cli']}/enviar", headers=_h(env),
                    json={"to": ["cobranza@cliente.mx"], "mensaje": "Buen día"})
    assert r.status_code == 200, r.text
    assert enviado["to"] == ["cobranza@cliente.mx"]
    assert "Estado de cuenta" in enviado["subject"]
    assert len(enviado["adj"]) == 1
    nombre, contenido, mime = enviado["adj"][0]
    assert nombre.endswith(".pdf") and mime == "application/pdf"
    assert contenido[:4] == b"%PDF"

    # Correo inválido: no llega al header To del SMTP.
    malo = client.post(f"/api/v1/cobranza/estado-cuenta/{env['cli']}/enviar", headers=_h(env),
                       json={"to": ["no-es-correo"]})
    assert malo.status_code == 422


def test_estado_cuenta_respeta_el_candado_por_cliente(client, env, auth_atado):
    """El usuario amarrado ve lo suyo; lo ajeno responde 404 por las tres
    puertas (JSON, PDF y correo), no solo por la del JSON."""
    h = _h(env)
    assert client.get(f"/api/v1/cobranza/estado-cuenta/{env['cli']}/pdf", headers=h).status_code == 200

    assert client.get(f"/api/v1/cobranza/estado-cuenta/{env['otro']}", headers=h).status_code == 404
    assert client.get(f"/api/v1/cobranza/estado-cuenta/{env['otro']}/pdf", headers=h).status_code == 404
    ajeno = client.post(f"/api/v1/cobranza/estado-cuenta/{env['otro']}/enviar", headers=h,
                        json={"to": ["quien@sea.mx"]})
    # 404 aunque el correo del tenant ni siquiera esté configurado: el candado va primero.
    assert ajeno.status_code == 404


# ── la tabla global de pendientes (rediseño de Cobranza, ticket 86bbyw5u2) ──


def test_facturas_pendientes_global_con_filtros(client, env, auth, fake_pac):
    """La tabla principal del rediseño: TODAS las PPD con saldo, filtrables
    por folio, y con el estado de pago derivado (PARCIAL cuando ya hay
    abonos). Las PUE y las saldadas no aparecen."""
    from decimal import Decimal as D

    h = _h(env)
    fid = _factura_ppd_timbrada(env, total=1000, dias_atras=5, folio=9001)
    _factura_ppd_timbrada(env, total=500, dias_atras=2, folio=9002)
    _factura_ppd_timbrada(env, total=300, dias_atras=1, metodo="PUE", folio=9003)

    r = client.get("/api/v1/cobranza/facturas-pendientes", headers=h)
    assert r.status_code == 200, r.text
    items = r.json()["items"]
    folios = {x["folio"] for x in items}
    assert {9001, 9002} <= folios and 9003 not in folios     # PUE fuera
    f1 = next(x for x in items if x["folio"] == 9001)
    assert f1["estado_pago"] == "PENDIENTE"
    assert x_ok(f1)

    # abono parcial → PARCIAL con el saldo actualizado
    rec = client.post("/api/v1/cobranza/recibos-pago", headers=h, json={
        "cliente_id": env["cli"], "fecha_pago": "2026-09-11T12:00:00Z",
        "forma_pago": "03", "monto": "400.00",
        "facturas": [{"factura_id": fid, "importe": "400.00"}],
    })
    assert rec.status_code == 201, rec.text
    # El saldo baja al TIMBRAR el REP (no al registrarlo): regla del módulo.
    rid = rec.json()["id"]
    t = client.post(f"/api/v1/cobranza/recibos-pago/{rid}/timbrar", headers=h)
    assert t.status_code == 200, t.text
    items = client.get("/api/v1/cobranza/facturas-pendientes", headers=h).json()["items"]
    f1 = next(x for x in items if x["folio"] == 9001)
    assert f1["estado_pago"] == "PARCIAL"
    assert D(str(f1["saldo_insoluto"])) == D("600")

    # filtro por folio (q)
    items = client.get("/api/v1/cobranza/facturas-pendientes?q=9002", headers=h).json()["items"]
    assert {x["folio"] for x in items} == {9002}


def x_ok(f):
    """Campos que la tabla necesita sí o sí."""
    return all(k in f for k in ("serie", "folio", "cliente_id", "fecha",
                                "vencimiento", "total", "saldo_insoluto", "estado_pago"))
