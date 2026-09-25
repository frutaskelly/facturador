"""Master de facturas: una fila por factura con remisión, OC, cobranza y NC."""
import io
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

import pytest
from sqlalchemy import text

from app.core.db import SessionLocal
from app.models import (
    Factura, NotaCredito, NotaCreditoFactura, OCRecibida, Remision, TimbradoIntento,
)

from tests.test_cobranza_api import (  # noqa: F401
    _factura_ppd_timbrada, _h, auth, auth_atado, env,
)

_URL = "/api/v1/reportes/master-facturas"


@pytest.fixture
def limpia(env):
    """Lo que el `env` de cobranza no purga y le estorba al borrar clientes."""
    yield
    db = SessionLocal()
    try:
        for tb in ("oc_recibidas", "nota_credito_facturas", "notas_credito", "lineas_remision", "remisiones"):
            db.execute(text(f"DELETE FROM {tb} WHERE tenant_id = :t"), {"t": env["tenant_id"]})
        db.commit()
    finally:
        db.close()


def _factura(env, **kw):
    db = SessionLocal()
    try:
        f = Factura(
            tenant_id=uuid.UUID(str(env["tenant_id"])), cliente_id=uuid.UUID(env["cli"]),
            forma_pago="03", metodo_pago="PPD", total=Decimal("100"), subtotal=Decimal("100"),
            fecha=datetime.now(timezone.utc) - timedelta(days=1), **kw,
        )
        db.add(f); db.commit()
        return f.id
    finally:
        db.close()


def _items(client, env, **params):
    r = client.get(_URL, params=params, headers=_h(env))
    assert r.status_code == 200, r.text
    return {f"{i['serie']}{i['folio']}": i for i in r.json()["items"]}, r.json()


def test_los_siete_estados(client, env, auth, limpia):
    _factura(env, serie="B", folio=1, estado="BORRADOR")
    fe = _factura(env, serie="B", folio=2, estado="BORRADOR")
    fp = _factura(env, serie="B", folio=3, estado="BORRADOR")
    _factura(env, serie="S", folio=4, estado="BORRADOR", origen="ESPEJO_SAE")
    _factura_ppd_timbrada(env, total=100, dias_atras=1, folio=5, serie="T")
    _factura_ppd_timbrada(env, total=100, dias_atras=1, folio=6, serie="T",
                          cancelacion_msj="En espera de aprobación")
    _factura_ppd_timbrada(env, total=100, dias_atras=1, folio=7, serie="T",
                          cancelacion_msj="No Cancelable")
    _factura(env, serie="C", folio=8, estado="CANCELADA", uuid=str(uuid.uuid4()), motivo_cancelacion="02")
    db = SessionLocal()
    try:
        t = uuid.UUID(str(env["tenant_id"]))
        db.add_all([
            TimbradoIntento(tenant_id=t, factura_id=fe, estado="ERROR", detalle="RFC del receptor inválido"),
            TimbradoIntento(tenant_id=t, factura_id=fp, estado="PENDIENTE"),
        ])
        db.commit()
    finally:
        db.close()

    items, d = _items(client, env)
    assert items["B1"]["estado"] == "BORRADOR"
    assert items["B2"]["estado"] == "ERROR_TIMBRADO"
    assert items["B2"]["estado_detalle"] == "RFC del receptor inválido"
    assert items["B3"]["estado"] == "TIMBRANDO"
    assert items["S4"]["estado"] == "ERROR_TIMBRADO"          # espejo sin UUID
    assert items["T5"]["estado"] == "TIMBRADA"
    assert items["T6"]["estado"] == "EN_CANCELACION"
    assert items["T7"]["estado"] == "NO_CANCELABLE"
    assert items["C8"]["estado"] == "CANCELADA"
    assert items["C8"]["motivo_cancelacion"].startswith("02")
    # Sin cobranza en lo que el SAT no da por vivo.
    assert items["B1"]["pagada"] is None and float(items["C8"]["saldo"]) == 0
    assert {e["clave"]: e["facturas"] for e in d["estados"]}["ERROR_TIMBRADO"] == 2

    solo, _ = _items(client, env, estado=["EN_CANCELACION", "ERROR_TIMBRADO"])
    assert set(solo) == {"B2", "S4", "T6"}
    assert client.get(_URL, params={"estado": "FOO"}, headers=_h(env)).status_code == 400


def test_cobranza_nc_remision_y_oc(client, env, auth, limpia):
    fid = uuid.UUID(_factura_ppd_timbrada(env, total=1000, dias_atras=40, folio=9, serie="ZX"))
    _factura_ppd_timbrada(env, total=500, dias_atras=2, metodo="PUE", folio=10, serie="ZX")
    db = SessionLocal()
    try:
        t = uuid.UUID(str(env["tenant_id"]))
        f = db.get(Factura, fid)
        f.saldo_insoluto = Decimal("600")         # $300 pagados + $100 de NC
        rem = Remision(tenant_id=t, folio_interno="RZX1", cliente_facturacion_id=f.cliente_id,
                       factura_id=fid, su_pedido="24736", fecha_remision=date.today())
        db.add(rem); db.flush()
        db.add(OCRecibida(tenant_id=t, canal="WHATSAPP", origen_externo=f"wa-{uuid.uuid4()}",
                          folio_externo="24736", remitente="5215550000000",
                          archivo_nombre="oc.pdf", archivo_url="https://drive.test/oc.pdf",
                          remision_id=rem.id))
        nc = NotaCredito(tenant_id=t, cliente_id=f.cliente_id, serie="NC", folio=1,
                         fecha=datetime.now(timezone.utc), total=Decimal("100"))
        db.add(nc); db.flush()
        db.add(NotaCreditoFactura(tenant_id=t, nota_id=nc.id, factura_id=fid, importe=Decimal("100")))
        db.commit()
    finally:
        db.close()

    desde = (date.today() - timedelta(days=60)).isoformat()
    items, _ = _items(client, env, desde=desde)
    x = items["ZX9"]
    assert x["remisiones"] == ["RZX1"] and x["su_pedido"] == "24736"
    assert x["oc_folio"] == "24736" and x["oc_url"] == "https://drive.test/oc.pdf"
    assert x["oc_canal"] == "WHATSAPP"
    assert float(x["nc_importe"]) == 100 and x["notas_credito"] == ["NC1"]
    assert float(x["pagado"]) == 300 and float(x["saldo"]) == 600 and x["pagada"] == "Parcial"
    assert x["dias_vencida"] == 10 and x["antiguedad"] == "1 mes"
    assert x["forma_pago_label"] == "99 Por definir"
    pue = items["ZX10"]
    assert pue["pagada"] == "Sí" and float(pue["pagado"]) == 500 and pue["dias_vencida"] is None

    r = client.get(f"{_URL}/xlsx", params={"desde": desde}, headers=_h(env))
    assert r.status_code == 200, r.text
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Master"]
    enc = [c.value for c in ws[2]]
    fila = next(row for row in ws.iter_rows(min_row=3) if row[1].value == 9)
    oc = fila[enc.index("OC original")]
    assert oc.value == "Ver OC" and oc.hyperlink.target == "https://drive.test/oc.pdf"
    pie = [c.value for c in ws[ws.max_row]]
    assert pie[0] == "TOTAL" and pie[enc.index("Total")].startswith("=SUBTOTAL(109,")
    assert wb["OC por factura"].max_row == 2


def test_candado_de_cliente(client, env, auth_atado, limpia):
    _factura_ppd_timbrada(env, total=100, dias_atras=1, folio=11, serie="K")
    r = client.get(_URL, params={"cliente_id": env["otro"]}, headers=_h(env))
    assert r.status_code == 404
