"""Grupos de cobranza: varias razones sociales en un solo estado de cuenta.

Dueño (2-oct-2026): EHMO quiere ver juntas sus razones sociales (EHMO, SUREÑA
y MAFAN) con la tabla por proyecto; otro cliente la querrá por serie, por
sucursal (estado) o por razón social. Un grupo dice qué entra, cómo se acomoda
la tabla y a quién se manda; aquí se arma el estado de cuenta, el correo, el
Excel (hoja Resumen + una hoja por fila) y el PDF.

Reglas:

1. **Un solo cálculo de saldos.** El estado de cuenta de cada razón social
   sale de `_armar_estado_cuenta`, el mismo del JSON, el PDF y el correo de
   hoy (saldo PPD, cancelaciones pedidas fuera, vencimiento con los días de
   crédito de SU razón social). El grupo solo filtra y junta.
2. **El proyecto es la fila de los reportes** (`ProyectoDeFactura.proyecto`,
   con «se reporta en»): IMSS BIENESTAR cae en HOSPITALES HIDALGO aquí igual
   que en Reportes.
3. **La sucursal sale del proyecto.** Lo que no cae en ningún proyecto usa el
   vínculo cliente↔sucursal de su serie. Un proyecto sin sucursal NO se
   adivina por la serie (ZMAFAN es Hidalgo, pero CDMX AZCAPOTZALCO no):
   sale «Sin sucursal» y se avisa.
4. **Lo que ves es lo que se manda.** El previo y el envío usan las mismas
   funciones de este módulo.
"""
from __future__ import annotations

import datetime as dt
import html as html_mod
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Callable, Optional
from uuid import UUID
from xml.sax.saxutils import escape as xml_escape

from sqlalchemy.orm import Session

from ..core.rbac import AuthContext
from ..models import (
    Cliente, ClienteSucursal, ClienteSucursalSerie, CobranzaConfig, CobranzaGrupo,
    CobranzaGrupoAlcance, Serie, Sucursal,
)
from .proyecto_de_factura import ProyectoDeFactura

ZERO = Decimal("0")
AGRUPAR_POR = ("PROYECTO", "SERIE", "SUCURSAL", "CLIENTE")
# Lo que dice «una hoja por …» en el correo y el Excel.
POR_TEXTO = {"PROYECTO": "proyecto", "SERIE": "serie", "SUCURSAL": "sucursal", "CLIENTE": "razón social"}
CUBETAS = (("por_vencer", "Por vencer"), ("d1_30", "1 a 30"), ("d31_60", "31 a 60"),
           ("d61_90", "61 a 90"), ("d90_mas", "Más de 90"))
SIN_SUCURSAL = "Sin sucursal"


# ─── La definición del grupo (guardado o recién capturado) ──────────────────

@dataclass
class Alcance:
    cliente_id: UUID
    completo: bool = True
    proyectos: set[UUID] = field(default_factory=set)
    series: set[str] = field(default_factory=set)

    def entra(self, doc: dict) -> bool:
        """¿Esta factura (un renglón del estado de cuenta) es del grupo?"""
        if self.completo:
            return True
        if doc.get("proyecto_id"):
            return UUID(doc["proyecto_id"]) in self.proyectos
        return (doc.get("serie") or "") in self.series

    def choca_con(self, otro: "Alcance") -> bool:
        """¿Cubren alguna factura en común? (para el aviso «también va en…»)."""
        if self.cliente_id != otro.cliente_id:
            return False
        if self.completo or otro.completo:
            return True
        return bool(self.proyectos & otro.proyectos) or bool(self.series & otro.series)


@dataclass
class Definicion:
    nombre: str
    agrupar_por: str = "PROYECTO"
    mostrar_antiguedad: bool = False
    correos: list[str] = field(default_factory=list)
    cc: list[str] = field(default_factory=list)
    pausado: bool = False
    alcance: list[Alcance] = field(default_factory=list)
    id: Optional[UUID] = None


def definicion_de(db: Session, grupo: CobranzaGrupo,
                  filas: Optional[list[CobranzaGrupoAlcance]] = None) -> Definicion:
    """La definición de un grupo guardado: sus filas de alcance por razón social."""
    if filas is None:
        filas = db.query(CobranzaGrupoAlcance).filter(
            CobranzaGrupoAlcance.tenant_id == grupo.tenant_id,
            CobranzaGrupoAlcance.grupo_id == grupo.id).all()
    por_cliente: dict[UUID, Alcance] = {}
    for f in filas:
        a = por_cliente.setdefault(f.cliente_id, Alcance(cliente_id=f.cliente_id, completo=False))
        if f.proyecto_id:
            a.proyectos.add(f.proyecto_id)
        elif f.serie:
            a.series.add(f.serie)
        else:
            a.completo = True
    for a in por_cliente.values():
        if a.completo:
            a.proyectos.clear()
            a.series.clear()
    return Definicion(
        id=grupo.id, nombre=grupo.nombre, agrupar_por=grupo.agrupar_por,
        mostrar_antiguedad=grupo.mostrar_antiguedad, correos=list(grupo.correos or []),
        cc=list(grupo.cc or []), pausado=grupo.pausado, alcance=list(por_cliente.values()),
    )


def filas_alcance(tenant_id, grupo_id, alcance: list[Alcance]) -> list[CobranzaGrupoAlcance]:
    """Las filas a guardar: una sin proyecto ni serie si es la razón social completa."""
    filas = []
    for a in alcance:
        if a.completo:
            filas.append(CobranzaGrupoAlcance(tenant_id=tenant_id, grupo_id=grupo_id, cliente_id=a.cliente_id))
            continue
        for p in sorted(a.proyectos, key=str):
            filas.append(CobranzaGrupoAlcance(tenant_id=tenant_id, grupo_id=grupo_id,
                                              cliente_id=a.cliente_id, proyecto_id=p))
        for s in sorted(a.series):
            filas.append(CobranzaGrupoAlcance(tenant_id=tenant_id, grupo_id=grupo_id,
                                              cliente_id=a.cliente_id, serie=s))
    return filas


# ─── Contexto compartido (un catálogo para varias razones sociales) ─────────

class Contexto:
    """Lo que se carga una vez por petición: el clasificador de proyectos, las
    sucursales y el estado de cuenta de cada razón social ya calculado (la
    lista de grupos pide el mismo cliente varias veces)."""

    def __init__(self, db: Session, ctx: AuthContext, corte: Optional[dt.date] = None):
        self.db = db
        self.ctx = ctx
        self.corte = corte or dt.datetime.now(dt.timezone.utc).date()
        self.clasificador = ProyectoDeFactura(db, ctx.tenant_id)
        self._estados: dict[UUID, dict] = {}
        self._clientes: dict[UUID, Cliente] = {}
        self.sucursales = {s.id: s.nombre for s in db.query(Sucursal).filter(
            Sucursal.tenant_id == ctx.tenant_id).all()}
        # (cliente, serie) → sucursal, por el vínculo: la serie default del
        # vínculo y su abanico. Solo para lo que no cae en ningún proyecto.
        self._sucursal_por_serie: dict[tuple[UUID, str], UUID] = {}
        codigos = {s.id: (s.codigo or "").strip().upper() for s in db.query(Serie).filter(
            Serie.tenant_id == ctx.tenant_id).all()}
        vinculos = db.query(ClienteSucursal).filter(ClienteSucursal.tenant_id == ctx.tenant_id).all()
        por_id = {v.id: v for v in vinculos}
        for v in vinculos:
            if v.serie_factura_id in codigos:
                self._sucursal_por_serie.setdefault((v.cliente_id, codigos[v.serie_factura_id]), v.sucursal_id)
        for cs in db.query(ClienteSucursalSerie).filter(ClienteSucursalSerie.tenant_id == ctx.tenant_id).all():
            v = por_id.get(cs.cliente_sucursal_id)
            if v is not None and cs.serie_id in codigos:
                self._sucursal_por_serie.setdefault((v.cliente_id, codigos[cs.serie_id]), v.sucursal_id)

    def cliente(self, cliente_id: UUID) -> Cliente:
        if cliente_id not in self._clientes:
            self._clientes[cliente_id] = self.db.query(Cliente).filter(
                Cliente.tenant_id == self.ctx.tenant_id, Cliente.id == cliente_id).one()
        return self._clientes[cliente_id]

    def estado(self, cliente_id: UUID) -> dict:
        """El estado de cuenta de la razón social (con su candado por cliente)."""
        from ..api.v1.cobranza import _armar_estado_cuenta

        if cliente_id not in self._estados:
            self._estados[cliente_id] = _armar_estado_cuenta(
                self.db, self.ctx, cliente_id, self.corte, clasificador=self.clasificador)
        return self._estados[cliente_id]

    def sucursal_de(self, cliente_id: UUID, doc: dict) -> Optional[UUID]:
        if doc.get("proyecto_id"):
            p = self.clasificador.por_id(UUID(doc["proyecto_id"]))
            return p.sucursal_id if p is not None else None
        return self._sucursal_por_serie.get((cliente_id, (doc.get("serie") or "").strip().upper()))


def nombre_corto(c: Cliente) -> str:
    return (c.nombre_corto or "").strip() or c.legal_name


# ─── El estado de cuenta del grupo ──────────────────────────────────────────

def _cubeta(dias_vencida: int) -> str:
    from ..api.v1.cobranza import _bucket
    return _bucket(dias_vencida)


def _nueva_fila() -> dict:
    return {"clientes": [], "series": [], "proyectos": [], "sucursales": [],
            "facturas": 0, "saldo": ZERO, "vencido": ZERO, "dias_max_vencida": 0,
            "antiguedad": {k: ZERO for k, _ in CUBETAS}, "docs": []}


def _agrega(lista: list, valor) -> None:
    if valor and valor not in lista:
        lista.append(valor)


def armar(cx: Contexto, d: Definicion, *, solo_vencidas: bool = False) -> dict:
    """El estado de cuenta del grupo: las filas de la tabla (cada una con sus
    facturas, que son la hoja de respaldo), los totales y los avisos."""
    if d.agrupar_por not in AGRUPAR_POR:
        raise ValueError(f"agrupar_por inválido: {d.agrupar_por}")
    # Las razones sociales en orden alfabético: EHMO, SUREÑA, MAFAN sale
    # así porque sus razones son GRUPO OPERADOR…, GRUPO SUREÑA, MEDIOS DE…
    alcance = sorted(d.alcance, key=lambda a: cx.cliente(a.cliente_id).legal_name or "")
    orden_cliente = {a.cliente_id: i for i, a in enumerate(alcance)}
    filas: dict[tuple, dict] = {}
    proyectos_sin_sucursal: set[str] = set()

    for a in alcance:
        cliente = cx.cliente(a.cliente_id)
        corto = nombre_corto(cliente)
        for doc in cx.estado(a.cliente_id)["facturas"]:
            if not a.entra(doc):
                continue
            if solo_vencidas and doc["dias_vencida"] <= 0:
                continue
            suc_id = cx.sucursal_de(a.cliente_id, doc)
            suc = cx.sucursales.get(suc_id) if suc_id else None
            if suc is None and doc.get("proyecto"):
                proyectos_sin_sucursal.add(doc["proyecto"])
            serie = doc.get("serie") or ""
            if d.agrupar_por == "PROYECTO":
                # Lo que no cae en ningún proyecto va al final de su razón social.
                clave = (orden_cliente[a.cliente_id], not doc.get("proyecto"), doc.get("proyecto") or "",
                         "" if doc.get("proyecto_id") else serie)
                hoja = doc.get("proyecto") or f"{corto} {serie}".strip()
            elif d.agrupar_por == "SERIE":
                clave = (orden_cliente[a.cliente_id], serie)
                hoja = serie or corto
            elif d.agrupar_por == "SUCURSAL":
                clave = (suc is None, suc or "")
                hoja = suc or SIN_SUCURSAL
            else:
                clave = (orden_cliente[a.cliente_id],)
                hoja = corto
            f = filas.get(clave)
            if f is None:
                f = filas[clave] = {**_nueva_fila(), "hoja": hoja}
            _agrega(f["clientes"], corto)
            _agrega(f["series"], serie)
            _agrega(f["proyectos"], doc.get("proyecto"))
            _agrega(f["sucursales"], suc)
            saldo = Decimal(doc["saldo_insoluto"])
            f["facturas"] += 1
            f["saldo"] += saldo
            if doc["dias_vencida"] > 0:
                f["vencido"] += saldo
            f["dias_max_vencida"] = max(f["dias_max_vencida"], doc["dias_vencida"])
            f["antiguedad"][_cubeta(doc["dias_vencida"])] += saldo
            f["docs"].append({**doc, "cliente_id": str(a.cliente_id), "cliente": corto,
                              "cliente_codigo": cliente.codigo, "sucursal": suc})

    lista = [filas[k] for k in sorted(filas)]
    for f in lista:
        f["docs"].sort(key=lambda x: (x["fecha"], x["serie"] or "", x["folio"]))
    total = {k: ZERO for k, _ in CUBETAS}
    for f in lista:
        for k in total:
            total[k] += f["antiguedad"][k]
    avisos = []
    if d.agrupar_por == "SUCURSAL" and proyectos_sin_sucursal:
        nombres = ", ".join(sorted(proyectos_sin_sucursal))
        avisos.append(f"Sin sucursal en Catálogo → Proyectos: {nombres}. Sus facturas salen en «{SIN_SUCURSAL}».")
    return {
        "nombre": d.nombre,
        "agrupar_por": d.agrupar_por,
        "mostrar_antiguedad": d.mostrar_antiguedad,
        "corte": cx.corte,
        "clientes": [{
            "cliente_id": str(a.cliente_id), "nombre": nombre_corto(cx.cliente(a.cliente_id)),
            "legal_name": cx.cliente(a.cliente_id).legal_name, "codigo": cx.cliente(a.cliente_id).codigo,
            "dias_credito": int(cx.cliente(a.cliente_id).dias_credito or 0),
        } for a in alcance],
        "filas": lista,
        "saldo_total": sum((f["saldo"] for f in lista), ZERO),
        "vencido_total": sum((f["vencido"] for f in lista), ZERO),
        "facturas": sum(f["facturas"] for f in lista),
        "dias_max_vencida": max((f["dias_max_vencida"] for f in lista), default=0),
        "antiguedad": total,
        "avisos": avisos,
    }


def resumen_sin_docs(datos: dict) -> dict:
    """Lo mismo sin las facturas de cada fila (para el JSON del previo)."""
    return {**datos, "filas": [{k: v for k, v in f.items() if k != "docs"} for f in datos["filas"]]}


# ─── Las columnas de la tabla (correo, Excel y PDF dicen lo mismo) ──────────

def _juntos(lista: list) -> str:
    return ", ".join(lista) if lista else "—"


def columnas(agrupar_por: str) -> list[tuple[str, Callable[[dict], str]]]:
    """Las columnas de texto de la tabla resumen; después van Saldo y Saldo vencido."""
    cliente = ("Cliente", lambda f: _juntos(f["clientes"]))
    if agrupar_por == "PROYECTO":
        return [cliente, ("Serie", lambda f: _juntos(f["series"])), ("Proyecto", lambda f: _juntos(f["proyectos"]))]
    if agrupar_por == "SERIE":
        return [cliente, ("Serie", lambda f: _juntos(f["series"])), ("Proyectos", lambda f: _juntos(f["proyectos"]))]
    if agrupar_por == "SUCURSAL":
        return [("Sucursal", lambda f: _juntos(f["sucursales"]) if f["sucursales"] else SIN_SUCURSAL),
                ("Clientes", lambda f: _juntos(f["clientes"])), ("Series", lambda f: _juntos(f["series"]))]
    return [cliente, ("Series", lambda f: _juntos(f["series"])), ("Proyectos", lambda f: _juntos(f["proyectos"]))]


def _pesos(v) -> str:
    return f"${Decimal(v or 0):,.2f}"


# ─── El correo ───────────────────────────────────────────────────────────────

_AZUL = "#305496"        # el azul de los reportes (reporte_pdf.AZUL)
_ZEBRA = "#EAF0FA"
_TOTAL = "#D9E1F2"
_REJILLA = "#BFBFBF"
_ROJO = "#C00000"


def asunto(cfg: CobranzaConfig, datos: dict) -> str:
    if (cfg.asunto or "").strip():
        return f"{cfg.asunto.strip()} · {datos['nombre']}"
    return f"Estado de cuenta {datos['nombre']} al {datos['corte']:%d/%m/%Y}"


def html_correo(cfg: CobranzaConfig, datos: dict, *, adjunta_excel: bool, encabezado: str = "") -> str:
    """El cuerpo del correo, con estilos en línea (los clientes de correo
    ignoran las hojas de estilo). Todo dato dinámico va con html.escape."""
    e = html_mod.escape
    mensaje = (cfg.mensaje or "").strip()
    partes = [encabezado] if encabezado else []
    partes += [f"<p>{e(p)}</p>" for p in mensaje.split("\n\n") if p.strip()]
    partes.append(f"<p>Estado de cuenta de <strong>{e(datos['nombre'])}</strong> al {datos['corte']:%d/%m/%Y}.</p>")
    n = datos["facturas"]
    if n == 0:
        partes.append("<p>No hay facturas por cobrar.</p>")
    else:
        vencido = datos["vencido_total"]
        partes.append(
            f"<p>Saldo total: <strong>{_pesos(datos['saldo_total'])}</strong> en {n:,} "
            f"{'factura' if n == 1 else 'facturas'} por cobrar"
            + (f", de los cuales <strong style=\"color:{_ROJO}\">{_pesos(vencido)}</strong> están vencidos"
               if vencido > 0 else "") + ".</p>")
        partes.append(_tabla_html(datos))
        if adjunta_excel:
            partes.append(
                f"<p style=\"font-size:12px;color:#6e6e73\">El detalle de cada factura va en el Excel adjunto: "
                f"una hoja por {POR_TEXTO[datos['agrupar_por']]}, más la hoja Resumen con esta misma tabla.</p>")
    return ("<div style=\"font-family:Arial,Helvetica,sans-serif;font-size:14px;color:#1d1d1f;line-height:1.45\">"
            + "".join(partes) + "</div>")


def _tabla_html(datos: dict) -> str:
    e = html_mod.escape
    cols = columnas(datos["agrupar_por"])
    cubetas = CUBETAS if datos["mostrar_antiguedad"] else ()
    th = (f"style=\"background:{_AZUL};color:#ffffff;text-align:left;padding:6px 10px;"
          "font-size:12px;font-weight:bold;white-space:nowrap\"")
    th_num = th.replace("text-align:left", "text-align:right")
    cab = "".join(f"<th {th}>{e(h)}</th>" for h, _ in cols)
    cab += f"<th {th_num}>Saldo</th><th {th_num}>Saldo vencido</th>"
    cab += "".join(f"<th {th_num}>{e(etq)}</th>" for _, etq in cubetas)
    cuerpo = []
    for i, f in enumerate(datos["filas"]):
        fondo = _ZEBRA if i % 2 else "#ffffff"
        td = f"style=\"padding:6px 10px;border-bottom:1px solid {_REJILLA};background:{fondo};vertical-align:top\""
        td_num = (f"style=\"padding:6px 10px;border-bottom:1px solid {_REJILLA};background:{fondo};"
                  "text-align:right;white-space:nowrap\"")
        td_venc = td_num.replace("text-align:right", f"text-align:right;color:{_ROJO}") if f["vencido"] > 0 else td_num
        celdas = "".join(f"<td {td}>{e(fn(f))}</td>" for _, fn in cols)
        celdas += f"<td {td_num}>{_pesos(f['saldo'])}</td><td {td_venc}>{_pesos(f['vencido'])}</td>"
        celdas += "".join(f"<td {td_num}>{_pesos(f['antiguedad'][k])}</td>" for k, _ in cubetas)
        cuerpo.append(f"<tr>{celdas}</tr>")
    tt = f"style=\"padding:6px 10px;background:{_TOTAL};font-weight:bold;border-top:2px solid {_AZUL}\""
    tt_num = tt.replace("font-weight:bold", "font-weight:bold;text-align:right;white-space:nowrap")
    total = f"<td {tt} colspan=\"{len(cols)}\">Total {e(datos['nombre'])}</td>"
    total += f"<td {tt_num}>{_pesos(datos['saldo_total'])}</td><td {tt_num}>{_pesos(datos['vencido_total'])}</td>"
    total += "".join(f"<td {tt_num}>{_pesos(datos['antiguedad'][k])}</td>" for k, _ in cubetas)
    return ("<table cellpadding=\"0\" cellspacing=\"0\" style=\"border-collapse:collapse;font-size:13px;"
            f"border:1px solid {_REJILLA};margin:8px 0\"><thead><tr>{cab}</tr></thead>"
            f"<tbody>{''.join(cuerpo)}<tr>{total}</tr></tbody></table>")


# ─── Archivos ────────────────────────────────────────────────────────────────

def nombre_archivo(datos: dict) -> str:
    """ASCII a propósito: el nombre viaja en Content-Disposition y un acento lo
    rompe en el navegador (mismo motivo que `_nombre_estado_cuenta`)."""
    import unicodedata
    base = unicodedata.normalize("NFKD", datos["nombre"]).encode("ascii", "ignore").decode()
    base = re.sub(r"[^A-Za-z0-9]+", "-", base).strip("-") or "grupo"
    return f"estado-cuenta-{base}-{datos['corte']:%Y%m%d}"


def nombres_hojas(datos: dict) -> list[str]:
    """Resumen + una hoja por fila, con las reglas de Excel: máximo 31
    caracteres, sin []:*?/\\ y sin repetir (dos razones sociales con la misma
    serie, p. ej. ZHGO de Balles y de Jubran)."""
    usados = {"resumen"}
    hojas = ["Resumen"]
    for f in datos["filas"]:
        base = re.sub(r"[\[\]:*?/\\]", " ", f["hoja"]).strip()[:31] or "Hoja"
        if base.lower() in usados and datos["agrupar_por"] == "SERIE" and f["clientes"]:
            base = f"{base} {f['clientes'][0]}"[:31]
        nombre, n = base, 2
        while nombre.lower() in usados:
            sufijo = f" ({n})"
            nombre = base[:31 - len(sufijo)] + sufijo
            n += 1
        usados.add(nombre.lower())
        hojas.append(nombre)
    return hojas


def xlsx(tenant, datos: dict) -> bytes:
    """Hoja Resumen (la tabla del correo) + una hoja por fila en el formato
    estilo SAE de `estado_cuenta_xlsx`."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    from .estado_cuenta_xlsx import MONEDA, escribir_tabla, fuentes

    FUENTES = fuentes()
    wb = Workbook()
    hojas = nombres_hojas(datos)
    ws = wb.active
    ws.title = hojas[0]

    def _fila(hoja, valores, font=FUENTES["normal"]):
        hoja.append(valores)
        for celda in hoja[hoja.max_row]:
            celda.font = font
        return hoja.max_row

    corte = f"{datos['corte']:%d/%m/%Y}"
    _fila(ws, [tenant.legal_name], FUENTES["titulo"])
    _fila(ws, [f"ESTADO DE CUENTA AL {corte} · {datos['nombre'].upper()}"], FUENTES["negrita"])
    for c in datos["clientes"]:
        _fila(ws, [f"{c['codigo'] or ''} {c['legal_name']} · {c['dias_credito']} días de crédito".strip()])
    _fila(ws, [])

    cols = columnas(datos["agrupar_por"])
    cubetas = CUBETAS if datos["mostrar_antiguedad"] else ()
    _fila(ws, [h.upper() for h, _ in cols] + ["SALDO", "SALDO VENCIDO", "FACTURAS"]
          + [etq.upper() for _, etq in cubetas], FUENTES["negrita"])
    primera = ws.max_row + 1
    for f in datos["filas"]:
        r = _fila(ws, [fn(f) for _, fn in cols] + [float(f["saldo"]), float(f["vencido"]), f["facturas"]]
                  + [float(f["antiguedad"][k]) for k, _ in cubetas],
                  FUENTES["vencida"] if f["vencido"] > 0 else FUENTES["normal"])
        for col in range(len(cols) + 1, len(cols) + 3):
            ws.cell(row=r, column=col).number_format = MONEDA
        for i in range(len(cubetas)):
            ws.cell(row=r, column=len(cols) + 4 + i).number_format = MONEDA
    ultima = ws.max_row
    total = [f"TOTAL {datos['nombre'].upper()}"] + [None] * (len(cols) - 1)
    num_cols = list(range(len(cols) + 1, len(cols) + 4 + len(cubetas)))
    for col in num_cols:
        letra = get_column_letter(col)
        total.append(f"=SUM({letra}{primera}:{letra}{ultima})" if ultima >= primera else 0)
    r = _fila(ws, total, FUENTES["negrita"])
    for col in num_cols:
        if col != len(cols) + 3:
            ws.cell(row=r, column=col).number_format = MONEDA
    anchos = [14, 14, 30][:len(cols)] if datos["agrupar_por"] != "SUCURSAL" else [16, 22, 26]
    for i, ancho in enumerate(anchos + [15, 15, 10] + [13] * len(cubetas), start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    for f, nombre in zip(datos["filas"], hojas[1:]):
        h = wb.create_sheet(nombre)
        _fila(h, [tenant.legal_name], FUENTES["titulo"])
        _fila(h, [f"ESTADO DE CUENTA AL {corte} · {f['hoja'].upper()}"], FUENTES["negrita"])
        _fila(h, [f"GRUPO {datos['nombre'].upper()}"], FUENTES["negrita"])
        _fila(h, [])
        escribir_tabla(h, f["docs"], cliente_de=lambda d: d["cliente"], etiqueta_total=f"TOTAL {f['hoja']}")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf(tenant, datos: dict) -> bytes:
    """Resumen (la tabla del correo), la antigüedad si el grupo la muestra y
    el detalle de cada fila con sus facturas."""
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer

    from .reporte_pdf import CELDA, construir, membrete, tabla_reporte

    def p(texto: str):
        return Paragraph(xml_escape(texto), CELDA)

    h4 = getSampleStyleSheet()["Heading4"]
    partes = membrete(tenant, "Estado de cuenta",
                      f"{datos['nombre']} · al {datos['corte']:%d/%m/%Y}")
    cols = columnas(datos["agrupar_por"])
    ancho_texto = 184 - 2 * 30
    anchos_texto = {3: [30, 34, ancho_texto - 64]}.get(len(cols), [ancho_texto / len(cols)] * len(cols))
    filas = [[p(fn(f)) for _, fn in cols] + [_pesos(f["saldo"]), _pesos(f["vencido"])] for f in datos["filas"]]
    filas.append([p(f"Total {datos['nombre']}")] + [""] * (len(cols) - 1)
                 + [_pesos(datos["saldo_total"]), _pesos(datos["vencido_total"])])
    partes.append(tabla_reporte(
        [h for h, _ in cols] + ["Saldo", "Saldo vencido"], filas,
        [w * mm for w in anchos_texto] + [30 * mm, 30 * mm],
        num_cols=(len(cols), len(cols) + 1), filas_totales=1))
    if datos["mostrar_antiguedad"]:
        partes.append(Spacer(1, 6 * mm))
        partes.append(Paragraph("Antigüedad de saldos", h4))
        partes.append(tabla_reporte(
            [etq for _, etq in CUBETAS], [[_pesos(datos["antiguedad"][k]) for k, _ in CUBETAS]],
            [36.8 * mm] * 5, num_cols=(0, 1, 2, 3, 4)))

    for f in datos["filas"]:
        partes.append(Spacer(1, 6 * mm))
        n = f["facturas"]
        partes.append(Paragraph(xml_escape(
            f"{f['hoja']} · {n} {'factura' if n == 1 else 'facturas'} · saldo {_pesos(f['saldo'])}"), h4))
        detalle = [[d["cliente"], f"{d['serie'] or ''}{d['folio']}", f"{d['fecha']:%d/%m/%Y}",
                    f"{d['vencimiento']:%d/%m/%Y}", str(d["dias_vencida"]) if d["dias_vencida"] > 0 else "",
                    _pesos(d["total"]), _pesos(d["saldo_insoluto"])] for d in f["docs"]]
        detalle.append(["", "", "", "", "Total", "", _pesos(f["saldo"])])
        partes.append(tabla_reporte(
            ["Cliente", "Factura", "Fecha", "Vence", "Días venc.", "Total", "Saldo"], detalle,
            [34 * mm, 30 * mm, 22 * mm, 22 * mm, 18 * mm, 29 * mm, 29 * mm],
            num_cols=(4, 5, 6), filas_totales=1))
    return construir("Estado de cuenta", partes)


# ─── El previo y los adjuntos (lo mismo que saldría en el correo) ───────────

def adjuntos(tenant, cfg: CobranzaConfig, datos: dict) -> list[tuple[str, bytes, str]]:
    """Los archivos del correo, según lo que diga Ajustes (PDF y/o Excel)."""
    nombre = nombre_archivo(datos)
    out = []
    if cfg.adjuntar_pdf:
        out.append((f"{nombre}.pdf", pdf(tenant, datos), "application/pdf"))
    if cfg.adjuntar_excel:
        out.append((f"{nombre}.xlsx", xlsx(tenant, datos),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
    return out


def previo(cx: Contexto, cfg: CobranzaConfig, d: Definicion) -> dict:
    """El correo tal como saldría hoy: destinatarios (con la copia de Ajustes y
    el escalamiento), asunto, cuerpo, los adjuntos que lleva y los avisos de
    por qué la cola semanal no lo mandaría."""
    from .cobranza_auto import _destinatarios, espejo_al_dia

    datos = armar(cx, d, solo_vencidas=not cfg.incluir_por_vencer)
    para, cc, escalado = _destinatarios(cfg, d, datos["dias_max_vencida"])
    avisos = list(datos["avisos"])
    if not para:
        avisos.insert(0, "Falta el correo: captura a quién se le manda (Para).")
    if d.pausado:
        avisos.append("El grupo está en pausa: la cobranza automática no lo manda.")
    if datos["facturas"] == 0:
        avisos.append("No hay facturas por cobrar con lo que tiene marcado el grupo.")
    elif datos["saldo_total"] < Decimal(cfg.saldo_minimo or 0):
        avisos.append(f"El saldo queda bajo el mínimo de Ajustes ({_pesos(cfg.saldo_minimo)}): "
                      "la cobranza automática no lo mandaría.")
    if not cfg.incluir_por_vencer:
        avisos.append("Solo van las facturas vencidas: en Ajustes está apagado «Incluir facturas por vencer».")
    ok, ultima = espejo_al_dia(cx.db, cx.ctx.tenant_id, cfg)
    if not ok:
        cuando = f"{ultima:%d/%m %H:%M} UTC" if ultima else "nunca"
        avisos.append(f"El espejo de SAE no está al día (última pasada buena: {cuando}): "
                      "los saldos pueden no traer los últimos pagos.")
    nombre = nombre_archivo(datos)
    hojas = nombres_hojas(datos)
    archivos = []
    if cfg.adjuntar_excel:
        n = len(hojas) - 1
        archivos.append({"tipo": "xlsx", "nombre": f"{nombre}.xlsx",
                         "detalle": f"Resumen + {n} {'hoja' if n == 1 else 'hojas'} por {POR_TEXTO[d.agrupar_por]}"})
    if cfg.adjuntar_pdf:
        archivos.append({"tipo": "pdf", "nombre": f"{nombre}.pdf",
                         "detalle": f"Resumen y detalle por {POR_TEXTO[d.agrupar_por]}"})
    return {
        "para": para, "cc": cc, "escalado": escalado,
        "asunto": asunto(cfg, datos),
        "html": html_correo(cfg, datos, adjunta_excel=cfg.adjuntar_excel),
        "adjuntos": archivos,
        "hojas": hojas if cfg.adjuntar_excel else [],
        "resumen": resumen_sin_docs(datos),
        "avisos": avisos,
    }
