"""El master de facturas: una fila por factura con todo lo que cuelga de ella.

Pedido del dueño (25-sep-2026): una sola sábana donde cada factura trae su
remisión, su OC (con la liga al documento original), su cobranza, sus notas de
crédito y su estado REAL. Nada de aquí se captura: todo se arma de lo que ya
guardan facturas, remisiones, la bandeja de OC, los REP y las NC del espejo.

Tres criterios que conviene no perder de vista:

  · El ESTADO del master no es la columna `facturas.estado`. BORRADOR mezcla
    tres cosas distintas (borrador de verdad, timbrado en vuelo, timbrado que
    falló sin UUID) y TIMBRADA otras tres (vigente, cancelación pedida al SAT,
    cancelación rechazada). El master las separa: son siete.
  · Pagado = total − saldo − NC, el mismo saldo del estado de cuenta, para que
    un número de aquí cuadre allá. Una PUE nace con saldo 0: sale pagada.
  · La OC llega por la cadena factura → remisión(es) → OC recibida. Solo ~1 de
    cada 5 facturas espejo tiene remisión ligada; en las demás la OC se queda
    en el `su_pedido` y la liga vacía.
"""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta
from decimal import Decimal

from sqlalchemy.orm import Session

from ..models import (
    Cliente, Factura, NotaCredito, NotaCreditoFactura, OCRecibida, Proyecto,
    ReciboPago, ReciboPagoFactura, Remision, Sucursal, TimbradoIntento, User,
)
from .espejo_cruce import extraer_semana

ZERO = Decimal("0")

# (clave, etiqueta) en el orden del ciclo de vida de una factura.
ESTADOS = (
    ("BORRADOR", "Borrador"),
    ("TIMBRANDO", "Timbrando…"),
    ("ERROR_TIMBRADO", "Error al timbrar (sin UUID)"),
    ("TIMBRADA", "Timbrada"),
    ("EN_CANCELACION", "En proceso de cancelación"),
    ("NO_CANCELABLE", "Cancelación rechazada"),
    ("CANCELADA", "Cancelada"),
)
ETIQUETA_ESTADO = dict(ESTADOS)
# Los que el SAT da por vivos: solo ellos llevan cobranza (pagado, saldo, vencimiento).
_CON_COBRANZA = {"TIMBRADA", "EN_CANCELACION", "NO_CANCELABLE"}

MOTIVO_CANCELACION = {
    "01": "01 Con errores con relación",
    "02": "02 Con errores sin relación",
    "03": "03 No se llevó a cabo la operación",
    "04": "04 Operación nominativa en factura global",
}

FORMA_PAGO = {
    "01": "Efectivo", "02": "Cheque nominativo", "03": "Transferencia electrónica",
    "04": "Tarjeta de crédito", "05": "Monedero electrónico", "06": "Dinero electrónico",
    "08": "Vales de despensa", "12": "Dación en pago", "13": "Pago por subrogación",
    "14": "Pago por consignación", "15": "Condonación", "17": "Compensación",
    "23": "Novación", "24": "Confusión", "25": "Remisión de deuda",
    "26": "Prescripción o caducidad", "27": "A satisfacción del acreedor",
    "28": "Tarjeta de débito", "29": "Tarjeta de servicios",
    "30": "Aplicación de anticipos", "31": "Intermediario pagos", "99": "Por definir",
}


def _cubeta(dias_vencida: int) -> str:
    if dias_vencida <= 0:
        return "Por vencer"
    if dias_vencida <= 30:
        return "1 mes"
    if dias_vencida <= 60:
        return "2 meses"
    if dias_vencida <= 90:
        return "3 meses"
    return "4+ meses"


def _dia(v) -> date | None:
    if v is None:
        return None
    return v.date() if isinstance(v, datetime) else v


def estado_master(f: Factura, intento: TimbradoIntento | None) -> tuple[str, str | None, datetime | None]:
    """(clave, detalle, fecha del hecho) del estado real de la factura."""
    if f.estado == "CANCELADA":
        return "CANCELADA", f.cancelacion_msj, f.fecha_cancelacion
    msj = (f.cancelacion_msj or "").strip()
    if f.estado == "TIMBRADA":
        if msj and "no cancelable" in msj.lower():
            return "NO_CANCELABLE", msj, None
        if msj:
            return "EN_CANCELACION", msj, None
        return "TIMBRADA", None, f.fecha_timbrado
    # BORRADOR: ¿alguien intentó timbrarla?
    if intento is not None and intento.estado == "PENDIENTE":
        return "TIMBRANDO", "El intento de timbrado quedó sin respuesta del PAC", intento.created_at
    if intento is not None and intento.estado == "ERROR":
        return "ERROR_TIMBRADO", intento.detalle, intento.created_at
    if f.origen != "NATIVA" and not f.uuid:
        # Espejo: SAE emitió el documento y el SAT nunca devolvió el UUID
        # (ZEHMOHOS 829/830). Se refleja BORRADOR por regla del dueño.
        return "ERROR_TIMBRADO", "SAE emitió el documento sin UUID del SAT", f.updated_at
    return "BORRADOR", None, None


def construir(
    db: Session, ctx, *, desde: date, hasta: date, cliente_id=None,
    estados: set[str] | None = None, plaza_de=None, proyecto_de_serie=None,
) -> list[dict]:
    """Las filas del master, ya en el orden del reporte (fecha desc, folio desc).

    `plaza_de(f, cliente_id)` y `proyecto_de_serie(f, nombre_cliente)` son los
    rescates de reportes.py para las facturas sin remisión ni proyecto ligados:
    la misma plaza y el mismo proyecto con que la cartera las reparte.
    """
    import sqlalchemy as sa

    q = (
        db.query(Factura, Cliente)
        .join(Cliente, Cliente.id == Factura.cliente_id)
        .filter(
            Factura.tenant_id == ctx.tenant_id,
            Factura.deleted_at.is_(None),
            sa.cast(Factura.fecha, sa.Date) >= desde,
            sa.cast(Factura.fecha, sa.Date) <= hasta,
        )
    )
    if ctx.cliente_scope:
        q = q.filter(Factura.cliente_id.in_(ctx.cliente_scope))
    if cliente_id is not None:
        q = q.filter(Factura.cliente_id == cliente_id)
    pares = q.order_by(Factura.fecha.desc(), Factura.serie, Factura.folio.desc()).all()
    if not pares:
        return []
    ids = [f.id for f, _ in pares]

    # ── Todo lo que cuelga, en una consulta por tabla ───────────────────────
    intentos: dict = {}
    for it in (db.query(TimbradoIntento)
               .filter(TimbradoIntento.factura_id.in_(ids))
               .order_by(TimbradoIntento.created_at.asc())):
        intentos[it.factura_id] = it          # el último gana

    remisiones: dict = {}
    for r in (db.query(Remision)
              .filter(Remision.factura_id.in_(ids), Remision.deleted_at.is_(None))
              .order_by(Remision.fecha_remision.asc(), Remision.folio_interno.asc())):
        remisiones.setdefault(r.factura_id, []).append(r)
    rem_ids = [r.id for rs in remisiones.values() for r in rs]

    ocs: dict = {}
    if rem_ids:
        for oc in (db.query(OCRecibida)
                   .filter(OCRecibida.tenant_id == ctx.tenant_id, OCRecibida.remision_id.in_(rem_ids))
                   .order_by(OCRecibida.recibida_at.asc())):
            ocs.setdefault(oc.remision_id, []).append(oc)

    proyecto_ids = {f.proyecto_id for f, _ in pares if f.proyecto_id}
    proyecto_ids |= {r.proyecto_id for rs in remisiones.values() for r in rs if r.proyecto_id}
    proyectos = {p.id: p for p in db.query(Proyecto).filter(Proyecto.id.in_(proyecto_ids))} if proyecto_ids else {}

    sucursal_ids = {r.sucursal_id for rs in remisiones.values() for r in rs if r.sucursal_id}
    sucursal_ids |= {p.sucursal_id for p in proyectos.values() if p.sucursal_id}
    sucursales = {s.id: s.nombre for s in db.query(Sucursal.id, Sucursal.nombre)
                  .filter(Sucursal.id.in_(sucursal_ids))} if sucursal_ids else {}

    reps: dict = {}
    for rf, rp in (db.query(ReciboPagoFactura, ReciboPago)
                   .join(ReciboPago, ReciboPago.id == ReciboPagoFactura.recibo_id)
                   .filter(ReciboPagoFactura.factura_id.in_(ids), ReciboPago.estado == "TIMBRADO")
                   .order_by(ReciboPago.fecha_pago.asc())):
        reps.setdefault(rf.factura_id, []).append((rf, rp))

    ncs: dict = {}
    for nf, nc in (db.query(NotaCreditoFactura, NotaCredito)
                   .join(NotaCredito, NotaCredito.id == NotaCreditoFactura.nota_id)
                   .filter(NotaCreditoFactura.factura_id.in_(ids), NotaCredito.estado != "CANCELADA")
                   .order_by(NotaCredito.fecha.asc())):
        ncs.setdefault(nf.factura_id, []).append((nf, nc))

    # Sustitución en los dos sentidos: por id (nueva → vieja) y por UUID (vieja → nueva).
    previas = {f.sustituye_a_factura_id for f, _ in pares if f.sustituye_a_factura_id}
    uuids_nuevos = {f.uuid_sustitucion for f, _ in pares if f.uuid_sustitucion}
    folio_por_id: dict = {}
    folio_por_uuid: dict = {}
    if previas or uuids_nuevos:
        for fid, serie, folio, uuid in (db.query(Factura.id, Factura.serie, Factura.folio, Factura.uuid)
                                        .filter(Factura.tenant_id == ctx.tenant_id,
                                                sa.or_(Factura.id.in_(previas or [None]),
                                                       Factura.uuid.in_(uuids_nuevos or [""])))):
            folio_por_id[fid] = f"{serie}{folio}"
            if uuid:
                folio_por_uuid[uuid] = f"{serie}{folio}"

    autores = {f.created_by for f, _ in pares if f.created_by}
    nombres = {u.id: (u.full_name or u.email) for u in
               db.query(User.id, User.full_name, User.email).filter(User.id.in_(autores))} if autores else {}

    hoy = datetime.now().date()
    filas: list[dict] = []
    for f, c in pares:
        if not ctx.cliente_permitido(c.id):
            continue
        estado, detalle, estado_fecha = estado_master(f, intentos.get(f.id))
        if estados and estado not in estados:
            continue

        rs = remisiones.get(f.id, [])
        primera = rs[0] if rs else None
        proyecto = proyectos.get(f.proyecto_id) or next(
            (proyectos[r.proyecto_id] for r in rs if r.proyecto_id in proyectos), None)
        sucursal = next((sucursales[r.sucursal_id] for r in rs if r.sucursal_id in sucursales), None)
        if sucursal is None and proyecto is not None and proyecto.sucursal_id in sucursales:
            sucursal = sucursales[proyecto.sucursal_id]
        if sucursal is None and plaza_de is not None:
            sucursal = plaza_de(f, c.id)
        nombre_proyecto = proyecto.nombre if proyecto else (
            proyecto_de_serie(f, c.legal_name) if proyecto_de_serie else None)
        if nombre_proyecto == c.legal_name:
            nombre_proyecto = None      # el rescate cae al cliente: no es un proyecto

        su_pedido = f.su_pedido or next((r.su_pedido for r in rs if r.su_pedido), None)
        oc_lista = [
            {
                "remision": r.folio_interno,
                "folio": oc.folio_externo,
                "url": oc.archivo_url,
                "archivo": oc.archivo_nombre,
                "canal": oc.canal,
                "remitente": oc.remitente,
                "recibida_at": oc.recibida_at,
                "fecha_entrega": oc.fecha_entrega,
                "punto_entrega": oc.punto_entrega,
                "cambio_at": oc.cambio_detectado_at if not oc.cambio_resuelto_at else None,
            }
            for r in rs for oc in ocs.get(r.id, [])
        ]
        oc_con_liga = next((o for o in oc_lista if o["url"]), None)

        total = Decimal(f.total or 0)
        nc_lista = ncs.get(f.id, [])
        nc_importe = sum((Decimal(nf.importe) for nf, _ in nc_lista), ZERO)
        rep_lista = reps.get(f.id, [])
        cobra = estado in _CON_COBRANZA
        saldo = Decimal(f.saldo_insoluto or 0) if cobra else ZERO
        pagado = max(total - saldo - nc_importe, ZERO) if cobra else ZERO
        if not cobra:
            pagada = None
        elif saldo <= Decimal("0.005"):
            pagada = "Sí"
        elif pagado > 0:
            pagada = "Parcial"
        else:
            pagada = "No"
        fecha = _dia(f.fecha)
        dias_credito = int(c.dias_credito or 0)
        vencimiento = fecha + timedelta(days=dias_credito) if cobra and fecha else None
        dias_vencida = (hoy - vencimiento).days if vencimiento and saldo > 0 else None
        ultimo = rep_lista[-1][1] if rep_lista else None

        filas.append({
            "id": str(f.id),
            # 1. Identificación
            "serie": f.serie, "folio": f.folio, "fecha": fecha,
            "fecha_timbrado": f.fecha_timbrado,
            "estado": estado, "estado_label": ETIQUETA_ESTADO[estado],
            "estado_detalle": detalle, "estado_fecha": estado_fecha,
            "uuid": f.uuid, "empresa": f.espejo_empresa,
            "origen": "SAE" if f.origen != "NATIVA" else "Facturador",
            "semana": extraer_semana(f.notas, su_pedido),
            # 2. Cliente y destino
            "cliente_id": str(c.id), "cliente": c.legal_name, "rfc": c.rfc,
            "cliente_codigo": c.codigo, "cliente_tipo": c.tipo,
            "proyecto": nombre_proyecto, "sucursal": sucursal,
            # 3. Origen comercial
            "remisiones": [r.folio_interno for r in rs],
            "su_pedido": su_pedido,
            "fecha_remision": primera.fecha_remision if primera else None,
            "fecha_entrega": primera.fecha_entrega if primera else None,
            "canal": primera.canal if primera else None,
            "factura_sae": next((r.factura_sae for r in rs if r.factura_sae), None),
            # 3b. Orden de compra
            "ocs": oc_lista,
            "oc_folio": ", ".join(dict.fromkeys(o["folio"] for o in oc_lista if o["folio"])) or None,
            "oc_url": oc_con_liga["url"] if oc_con_liga else None,
            "oc_archivo": oc_con_liga["archivo"] if oc_con_liga else None,
            "oc_canal": oc_lista[0]["canal"] if oc_lista else None,
            "oc_remitente": oc_lista[0]["remitente"] if oc_lista else None,
            "oc_recibida_at": oc_lista[0]["recibida_at"] if oc_lista else None,
            "oc_fecha_entrega": oc_lista[0]["fecha_entrega"] if oc_lista else None,
            "oc_punto_entrega": oc_lista[0]["punto_entrega"] if oc_lista else None,
            "oc_cambio_at": next((o["cambio_at"] for o in oc_lista if o["cambio_at"]), None),
            # 4. Importes
            "subtotal": f.subtotal, "descuento": f.descuento,
            "iva": f.iva_trasladado, "ieps": f.ieps_trasladado,
            "ret_iva": f.ret_iva, "ret_isr": f.ret_isr,
            "total": total, "moneda": f.moneda,
            # 5. Fiscal
            "forma_pago": f.forma_pago,
            "forma_pago_label": f"{f.forma_pago} {FORMA_PAGO.get(f.forma_pago, '')}".strip(),
            "metodo_pago": f.metodo_pago, "uso_cfdi": f.uso_cfdi,
            "tipo_comprobante": f.tipo_comprobante,
            # 6. Cobranza
            "pagado": pagado, "saldo": saldo, "pagada": pagada,
            "ultimo_pago": _dia(ultimo.fecha_pago) if ultimo else None,
            "parcialidades": max((rf.num_parcialidad for rf, _ in rep_lista), default=0) or None,
            "reps": [f"{rp.serie}{rp.folio}" for _, rp in rep_lista],
            "banco": ultimo.banco if ultimo else None,
            "referencia": ultimo.num_operacion if ultimo else None,
            "dias_credito": dias_credito,
            "vencimiento": vencimiento,
            "dias_vencida": dias_vencida,
            "antiguedad": _cubeta(dias_vencida) if dias_vencida is not None else None,
            # 7. Notas de crédito
            "notas_credito": [f"{nc.serie}{nc.folio}" for _, nc in nc_lista],
            "nc_importe": nc_importe,
            "total_neto": total - nc_importe,
            # 8. Cancelación y sustitución
            "fecha_cancelacion": f.fecha_cancelacion,
            "motivo_cancelacion": MOTIVO_CANCELACION.get(f.motivo_cancelacion or "", f.motivo_cancelacion),
            "sustituye_a": folio_por_id.get(f.sustituye_a_factura_id),
            "sustituida_por": folio_por_uuid.get(f.uuid_sustitucion) or f.uuid_sustitucion,
            # 9. Otros
            "notas": f.notas,
            "creada_por": nombres.get(f.created_by),
        })
    return filas


# ── Excel ────────────────────────────────────────────────────────────────────
#
# (encabezado, clave o función, formato, ancho). La fila de totales usa
# SUBTOTAL(109): si el usuario filtra en Excel, el pie suma solo lo visible.

def _lista(clave):
    return lambda d: ", ".join(d[clave]) or None


_FECHA = "DD/MM/YYYY"
_MONEDA = "#,##0.00"
_COLUMNAS = [
    ("Serie", "serie", None, 9),
    ("Folio", "folio", "0", 8),
    ("Fecha", "fecha", _FECHA, 11),
    ("Estado", "estado_label", None, 24),
    ("Detalle del estado", "estado_detalle", None, 30),
    ("Semana", "semana", "0", 8),
    ("Cliente", "cliente", None, 32),
    ("RFC", "rfc", None, 14),
    ("Proyecto", "proyecto", None, 26),
    ("Sucursal", "sucursal", None, 16),
    ("Remisión(es)", _lista("remisiones"), None, 18),
    ("Su pedido", "su_pedido", None, 16),
    ("Folio OC", "oc_folio", None, 16),
    ("OC original", None, None, 12),            # hipervínculo, se llena aparte
    ("Canal OC", "oc_canal", None, 11),
    ("Remitente OC", "oc_remitente", None, 22),
    ("OC recibida", "oc_recibida_at", _FECHA, 11),
    ("OC cambió", "oc_cambio_at", _FECHA, 11),
    ("Fecha remisión", "fecha_remision", _FECHA, 11),
    ("Fecha entrega", "fecha_entrega", _FECHA, 11),
    ("Subtotal", "subtotal", _MONEDA, 13),
    ("Descuento", "descuento", _MONEDA, 11),
    ("IVA", "iva", _MONEDA, 12),
    ("IEPS", "ieps", _MONEDA, 11),
    ("Ret. IVA", "ret_iva", _MONEDA, 10),
    ("Ret. ISR", "ret_isr", _MONEDA, 10),
    ("Total", "total", _MONEDA, 13),
    ("Moneda", "moneda", None, 7),
    ("Forma de pago", "forma_pago_label", None, 24),
    ("Método de pago", "metodo_pago", None, 8),
    ("Uso CFDI", "uso_cfdi", None, 8),
    ("Pagado", "pagado", _MONEDA, 13),
    ("Saldo", "saldo", _MONEDA, 13),
    ("¿Pagada?", "pagada", None, 9),
    ("Último pago", "ultimo_pago", _FECHA, 11),
    ("Parcialidades", "parcialidades", "0", 8),
    ("REP", _lista("reps"), None, 16),
    ("Banco", "banco", None, 16),
    ("Referencia", "referencia", None, 16),
    ("Días crédito", "dias_credito", "0", 8),
    ("Vencimiento", "vencimiento", _FECHA, 11),
    ("Días vencida", "dias_vencida", "0", 8),
    ("Antigüedad", "antiguedad", None, 11),
    ("Notas de crédito", _lista("notas_credito"), None, 16),
    ("Importe NC", "nc_importe", _MONEDA, 12),
    ("Total neto", "total_neto", _MONEDA, 13),
    ("Fecha cancelación", "fecha_cancelacion", _FECHA, 11),
    ("Motivo cancelación", "motivo_cancelacion", None, 22),
    ("Sustituye a", "sustituye_a", None, 12),
    ("Sustituida por", "sustituida_por", None, 12),
    ("UUID", "uuid", None, 38),
    ("Empresa SAE", "empresa", None, 8),
    ("Origen", "origen", None, 10),
    ("Fecha timbrado", "fecha_timbrado", _FECHA, 11),
    ("Creada por", "creada_por", None, 20),
    ("Nota", "notas", None, 40),
]
_SUMABLES = {"subtotal", "descuento", "iva", "ieps", "ret_iva", "ret_isr", "total",
             "pagado", "saldo", "nc_importe", "total_neto"}


def _celda(v):
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime):
        return v.replace(tzinfo=None).date() if v.tzinfo else v.date()
    return v


def generar_xlsx(filas: list[dict], *, titulo: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    negrita = Font(bold=True)
    liga = Font(color="FF0563C1", underline="single")
    cabeza = PatternFill("solid", fgColor="FFE7E6E6")

    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([c[0] for c in _COLUMNAS])
    for celda in ws[2]:
        celda.font = negrita
        celda.fill = cabeza
    col_oc = next(i for i, c in enumerate(_COLUMNAS, start=1) if c[0] == "OC original")

    primera = 3
    r = primera - 1            # a mano: `ws.max_row` recorre la hoja y vuelve cuadrático el armado
    for d in filas:
        valores = []
        for _, clave, _, _ in _COLUMNAS:
            if clave is None:
                valores.append(None)
            elif callable(clave):
                valores.append(clave(d))
            else:
                valores.append(_celda(d.get(clave)))
        ws.append(valores)
        r += 1
        for i, (_, _, fmt, _) in enumerate(_COLUMNAS, start=1):
            if fmt:
                ws.cell(row=r, column=i).number_format = fmt
        if d["oc_url"]:
            n = len([o for o in d["ocs"] if o["url"]])
            c = ws.cell(row=r, column=col_oc, value="Ver OC" if n == 1 else f"Ver OC ({n})")
            c.hyperlink = d["oc_url"]
            c.font = liga
        elif d["ocs"] or d["su_pedido"]:
            ws.cell(row=r, column=col_oc, value="sin archivo")
    ultima = r

    # Pie: SUBTOTAL(109) respeta el autofiltro — filtra y el total se mueve.
    pie = ["TOTAL"] + [None] * (len(_COLUMNAS) - 1)
    ws.append(pie)
    r += 1
    ws.cell(row=r, column=1).font = negrita
    for i, (_, clave, fmt, _) in enumerate(_COLUMNAS, start=1):
        if clave in _SUMABLES:
            letra = get_column_letter(i)
            c = ws.cell(row=r, column=i,
                        value=f"=SUBTOTAL(109,{letra}{primera}:{letra}{max(ultima, primera)})")
            c.font = negrita
            c.number_format = fmt

    for i, (_, _, _, ancho) in enumerate(_COLUMNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(_COLUMNAS))}{max(ultima, 2)}"

    # ── Hoja 2: una fila por OC, con su liga (una celda no lleva dos) ───────
    oc = wb.create_sheet("OC por factura")
    enc = ["Factura", "Cliente", "Remisión", "Folio OC", "OC original", "Archivo",
           "Canal", "Remitente", "Recibida", "Fecha entrega", "Punto de entrega"]
    oc.append(enc)
    for celda in oc[1]:
        celda.font = negrita
        celda.fill = cabeza
    r = 1
    for d in filas:
        for o in d["ocs"]:
            oc.append([f"{d['serie']}{d['folio']}", d["cliente"], o["remision"], o["folio"],
                       "Ver OC" if o["url"] else "sin archivo", o["archivo"], o["canal"],
                       o["remitente"], _celda(o["recibida_at"]), _celda(o["fecha_entrega"]),
                       o["punto_entrega"]])
            r += 1
            if o["url"]:
                oc.cell(row=r, column=5).hyperlink = o["url"]
                oc.cell(row=r, column=5).font = liga
            oc.cell(row=r, column=9).number_format = _FECHA
            oc.cell(row=r, column=10).number_format = _FECHA
    for i, ancho in enumerate((12, 32, 16, 16, 12, 28, 10, 24, 11, 11, 28), start=1):
        oc.column_dimensions[get_column_letter(i)].width = ancho
    oc.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
