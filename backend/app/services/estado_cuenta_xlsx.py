"""El estado de cuenta en Excel, con el layout del que SAE le manda al cliente.

Pedido del dueño (11-sep-2026): EHMO recibe hoy su estado de cuenta como Excel
sacado de SAE (encabezado con crédito, columnas CLIENTE/TIPO/CONCEPTO/PROYECTO/
FACTURA/FECHA APLIC/SUBTOTAL/ABONOS/SALDOS, marca VENCIDO y renglón TOTAL), y
además le anotan a mano la SEMANA de cada entrega. Este módulo reproduce ese
formato con los datos combinados del Facturador y del espejo SAE — la semana
sale sola de las observaciones/su_pedido, ya no se captura.

Columnas del Excel (mismas que el de SAE, más SEM y ESTATUS con nombre):

    SEM | CLIENTE | TIPO | CONCEPTO | PROYECTO | FACTURA | FECHA APLIC |
    SUBTOTAL | ABONOS | SALDOS | ESTATUS

- ABONOS = total - saldo (el espejo trae el saldo de SAE, no cada pago).
- ESTATUS = "VENCIDO" cuando la fecha + días de crédito quedó atrás del corte.
- El renglón TOTAL suma con fórmulas =SUM(): el Excel sigue cuadrando si el
  cliente borra renglones ya pagados.
"""
from __future__ import annotations

import io
from decimal import Decimal

from ..models import Cliente, Tenant

HDR = ["SEM", "CLIENTE", "TIPO", "CONCEPTO", "PROYECTO", "FACTURA",
       "FECHA APLIC", "SUBTOTAL", "ABONOS", "SALDOS", "ESTATUS"]
_ANCHOS = (9, 10, 9, 11, 18, 18, 12, 13, 13, 13, 10)
_MONEDA = "#,##0.00"


def _domicilio(cliente: Cliente) -> list[str]:
    """Las líneas de dirección que existan, en el orden del membrete SAE."""
    dom = cliente.domicilio_fiscal or {}
    linea1 = " ".join(v for v in (dom.get("calle"), dom.get("colonia")) if v)
    linea2 = " ".join(v for v in (dom.get("cp"), dom.get("ciudad"), dom.get("estado")) if v)
    return [ln for ln in (linea1, linea2) if ln]


def generar(tenant: Tenant, cliente: Cliente, datos: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    normal = Font(name="Arial", size=10)
    negrita = Font(name="Arial", size=10, bold=True)
    titulo = Font(name="Arial", size=12, bold=True)
    vencida = Font(name="Arial", size=10, color="FFC00000")

    wb = Workbook()
    ws = wb.active
    ws.title = datos["serie"] or "Estado de cuenta"

    def _fila(valores, font=normal):
        ws.append(valores)
        for celda in ws[ws.max_row]:
            celda.font = font
        return ws.max_row

    # ── Membrete: emisor, corte y datos del cliente con su crédito ──────────
    _fila([tenant.legal_name], titulo)
    subtitulo = f"ESTADO DE CUENTA AL {datos['corte']:%d/%m/%Y}"
    if datos["serie"]:
        subtitulo += f" · SERIE {datos['serie']}"
    _fila([subtitulo], negrita)
    _fila([])
    _fila([cliente.legal_name, None, None, None, None, f"CLIENTE {cliente.codigo or ''}".strip()], negrita)
    limite = Decimal(cliente.limite_credito or 0)
    disponible = max(limite - Decimal(datos["saldo_total"]), Decimal("0"))
    derecha = [
        ("Días de crédito:", datos["dias_credito"]),
        ("Límite de crédito:", float(limite)),
        ("Saldo disponible:", float(disponible)),
        ("Moneda: Pesos", None),
    ]
    izquierda = _domicilio(cliente) + [cliente.rfc or ""]
    for i in range(max(len(izquierda), len(derecha))):
        fila = [izquierda[i] if i < len(izquierda) else None, None, None, None, None]
        fila += list(derecha[i]) if i < len(derecha) else [None, None]
        r = _fila(fila)
        if i < len(derecha) and derecha[i][1] is not None and i > 0:
            ws.cell(row=r, column=7).number_format = _MONEDA
    _fila([])

    # ── La tabla ────────────────────────────────────────────────────────────
    _fila(HDR, negrita)
    primera = ws.max_row + 1
    for d in datos["facturas"]:
        total = Decimal(d["total"])
        saldo = Decimal(d["saldo_insoluto"])
        abonos = total - saldo
        vencido = d["dias_vencida"] > 0
        r = _fila([
            f"SEM {d['semana']}" if d["semana"] else None,
            cliente.codigo,
            "Matriz",
            "Factura",
            d["proyecto"],
            f"{d['serie']} {d['folio']}",
            d["fecha"],
            float(total),
            float(abonos) if abonos > 0 else None,
            float(saldo),
            "VENCIDO" if vencido else None,
        ], vencida if vencido else normal)
        ws.cell(row=r, column=7).number_format = "DD/MM/YYYY"
        for col in (8, 9, 10):
            ws.cell(row=r, column=col).number_format = _MONEDA
    ultima = ws.max_row

    # ── TOTAL (estilo "TOTAL HOSPITALES" de SAE: proyecto único, o la serie) ─
    proyectos = {d["proyecto"] for d in datos["facturas"] if d["proyecto"]}
    etiqueta = "TOTAL"
    if len(proyectos) == 1:
        etiqueta = f"TOTAL {next(iter(proyectos))}"
    elif datos["serie"]:
        etiqueta = f"TOTAL {datos['serie']}"
    fila_total = [None, None, None, None, None, etiqueta, None]
    if ultima >= primera:
        for col in (8, 9, 10):
            letra = get_column_letter(col)
            fila_total.append(f"=SUM({letra}{primera}:{letra}{ultima})")
    else:
        fila_total += [0, 0, 0]
    r = _fila(fila_total, negrita)
    for col in (8, 9, 10):
        ws.cell(row=r, column=col).number_format = _MONEDA

    for i, ancho in enumerate(_ANCHOS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = f"A{primera}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
